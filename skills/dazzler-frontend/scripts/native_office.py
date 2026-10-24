"""Inspect editable Office structure or add a data-linked PPTX chart. No installation or rendering."""

import argparse
import hashlib
import io
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile


def read_package(path):
    with Path(path).open("rb") as stream:
        data = stream.read(32_000_001)
    if len(data) > 32_000_000:
        raise ValueError("Office input exceeds 32 MB")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = z.infolist()
        if len(entries) > 4096 or sum(e.file_size for e in entries) > 64_000_000:
            raise ValueError("Expanded Office package exceeds limits")
        result = {}
        for entry in entries:
            if entry.filename.endswith(".xml"):
                raw = z.read(entry)
                if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                    raise ValueError("XML declarations are not supported")
                result[entry.filename] = ET.fromstring(raw)
    return data, result


def inspect(path):
    data, parts = read_package(path)
    main = [
        v
        for k, v in parts.items()
        if k == "word/document.xml"
        or (k.startswith("ppt/slides/slide") and k.endswith(".xml"))
    ]
    if not main:
        raise ValueError("Expected a DOCX or PPTX package")
    fonts = set()
    for tree in main:
        for e in tree.iter():
            for key, value in e.attrib.items():
                if key.split("}")[-1] in [
                    "typeface",
                    "ascii",
                    "hAnsi",
                    "eastAsia",
                    "cs",
                ]:
                    fonts.add(value)
    return {
        "sha256": hashlib.sha256(data).hexdigest(),
        "contentParts": len(main),
        "textRuns": sum(
            e.tag.split("}")[-1] == "t" for tree in main for e in tree.iter()
        ),
        "nativeCharts": sum(
            k.startswith(("word/charts/chart", "ppt/charts/chart")) for k in parts
        ),
        "declaredFonts": sorted(fonts),
        "fontVerification": "unknown-until-rendered",
        "reviewStatus": "structure-only",
        "next": "Compare actual PDF font identities and every rendered page; test a separate edit copy.",
    }


def validate_chart(data):
    if not isinstance(data, dict) or set(data) != {"categories", "series"}:
        raise ValueError("Chart needs categories and series")
    cats, series = data["categories"], data["series"]
    if (
        not isinstance(cats, list)
        or not 1 <= len(cats) <= 100
        or any(not isinstance(c, str) or not c.strip() or len(c) > 300 for c in cats)
    ):
        raise ValueError("Use 1–100 nonempty category labels")
    if not isinstance(series, list) or not 1 <= len(series) <= 12:
        raise ValueError("Use 1–12 series")
    names = set()
    for s in series:
        if (
            not isinstance(s, dict)
            or set(s) != {"name", "values"}
            or not isinstance(s["name"], str)
            or not s["name"].strip()
            or len(s["name"]) > 300
            or s["name"] in names
        ):
            raise ValueError("Series names must be unique and nonempty")
        names.add(s["name"])
        if not isinstance(s["values"], list) or len(s["values"]) != len(cats):
            raise ValueError("Every category needs a value or explicit null")
        if any(
            v is not None
            and (
                isinstance(v, bool)
                or not isinstance(v, (int, float))
                or not math.isfinite(v)
                or abs(v) > 1e15
            )
            for v in s["values"]
        ):
            raise ValueError("Use bounded finite numbers or null; missing is not zero")
    return data


def add_chart(source, destination, data, slide, box, kind):
    validate_chart(data)
    read_package(source)
    if (
        kind not in ["column", "line"]
        or len(box) != 4
        or any(not math.isfinite(v) or v < 0 for v in box)
        or not box[2]
        or not box[3]
    ):
        raise ValueError("Invalid chart type or box")
    # Optional authoring library must already be available in the selected host runtime.
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation(source)
    x, y, w, h = [Inches(v) for v in box]
    if (
        not 1 <= slide <= len(deck.slides)
        or x + w > deck.slide_width
        or y + h > deck.slide_height
    ):
        raise ValueError("Chart must fit on an existing slide")
    chart_data = CategoryChartData()
    chart_data.categories = data["categories"]
    for series in data["series"]:
        chart_data.add_series(series["name"], series["values"])
    shape = deck.slides[slide - 1].shapes.add_chart(
        (
            XL_CHART_TYPE.COLUMN_CLUSTERED
            if kind == "column"
            else XL_CHART_TYPE.LINE_MARKERS
        ),
        x,
        y,
        w,
        h,
        chart_data,
    )
    blanks = shape.chart._chartSpace.xpath(".//c:dispBlanksAs")
    if len(blanks) != 1:
        raise ValueError("Unsupported chart serialization: missing blank-value policy")
    blanks[0].set("val", "gap")
    if kind == "column":
        shape.chart.value_axis.minimum_scale = (
            0
            if all(v is None or v >= 0 for s in data["series"] for v in s["values"])
            else None
        )
    stream = io.BytesIO()
    deck.save(stream)
    # Never overwrite the source or an existing artifact.
    with Path(destination).open("xb") as out:
        out.write(stream.getvalue())
    return {
        "nativeChart": True,
        "renderReview": "required",
        "sourceData": data,
        "note": "Style the native chart for the brief; verify labels, units, missing gaps and embedded workbook edits.",
    }


def main():
    p = argparse.ArgumentParser(description=__doc__)
    subs = p.add_subparsers(dest="command", required=True)
    check = subs.add_parser("inspect")
    check.add_argument("file", type=Path)
    chart = subs.add_parser("chart")
    chart.add_argument("file", type=Path)
    chart.add_argument("--out", required=True, type=Path)
    chart.add_argument("--data", required=True, type=Path)
    chart.add_argument("--slide", type=int, required=True)
    chart.add_argument(
        "--box",
        nargs=4,
        type=float,
        required=True,
        metavar=("X", "Y", "WIDTH", "HEIGHT"),
    )
    chart.add_argument("--kind", choices=["column", "line"], default="column")
    a = p.parse_args()
    if a.command == "inspect":
        result = inspect(a.file)
    else:
        if a.data.stat().st_size > 100_000:
            raise ValueError("Chart data exceeds 100 KB")
        result = add_chart(
            a.file,
            a.out,
            json.loads(a.data.read_text(encoding="utf-8")),
            a.slide,
            a.box,
            a.kind,
        )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
