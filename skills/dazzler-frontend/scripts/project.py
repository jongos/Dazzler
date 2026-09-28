"""Dazzler brand import, licensed assets, reversible changes and document exporters. Stdlib core."""
import argparse
from collections import Counter, defaultdict
import hashlib
import html
import json
from pathlib import Path
import re
import shutil

SKILL = Path(__file__).resolve().parents[1]


def read(path):
    with Path(path).open('rb') as handle:
        data=handle.read(8_000_001)
    if len(data)>8_000_000:raise ValueError('JSON input exceeds 8 MB')
    return json.loads(data.decode('utf-8'))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def inside(root, relative):
    p = (Path(root).resolve()/relative).resolve()
    if not p.is_relative_to(Path(root).resolve()) or p == Path(root).resolve():
        raise ValueError('Path must stay inside project root')
    return p


def brand(source):
    """Inventory CSS evidence; do not pretend a regex resolves the CSS cascade."""
    source = Path(source).resolve()
    files = [source] if source.is_file() else sorted(source.rglob('*.css'))
    rows = []; values = defaultdict(list)
    for p in files:
        if any(x in p.parts for x in ('.git','node_modules','dist')): continue
        css = re.sub(r'/\*.*?\*/', lambda m:'\n'*m.group().count('\n'), p.read_text(encoding='utf-8'), flags=re.S)
        for m in re.finditer(r'(--[\w-]+|color|background-color|font-family|font-size|border-radius|gap|padding|margin)\s*:\s*([^;{}]+)', css):
            key, value = m.groups(); value=value.strip()
            row={'file':str(p),'line':css[:m.start()].count('\n')+1,'property':key,'value':value}
            rows.append(row);values[key].append(row)
    conflicts=[{'property':key,'values':sorted({x['value'] for x in rs}),'sources':rs}
               for key,rs in values.items() if key.startswith('--') and len({x['value'] for x in rs})>1]
    families=Counter(x['value'] for x in rows if x['property']=='font-family')
    colors=Counter(x['value'] for x in rows if re.fullmatch(r'#[0-9a-fA-F]{6}',x['value']))
    return {'schemaVersion':1,'source':str(source),'observations':rows,'conflicts':conflicts,
            'candidates':{'fonts':families.most_common(),'colors':colors.most_common()},
            'locks':{},'status':'review-required' if conflicts else 'observed',
            'limitations':['Static declarations only: media queries, selectors, inheritance and variable resolution require rendered inspection.','Frequency is evidence of usage, not a brand authority decision. Confirm against existing brand rules before locking.']}


def assets(kind, mood, destination, ids=None):
    catalog=read(SKILL/'references/asset-catalog.json')
    selected=[x for x in catalog['assets'] if x['kind']==kind and (not ids or x['id'] in ids)]
    selected.sort(key=lambda x:(-sum(word in x['tags'] for word in mood.split()),x['id']))
    if not selected: raise ValueError('No matching assets')
    if ids and set(ids)-{x['id'] for x in selected}: raise ValueError('Unknown requested asset')
    for item in selected:
        source=inside(SKILL,item['path'])
        if digest(source)!=item['sha256']:raise ValueError('Asset hash mismatch')
    target=Path(destination);target.mkdir(parents=True,exist_ok=False)
    for item in selected:
        source=inside(SKILL,item['path'])
        shutil.copy2(source,target/source.name)
    shutil.copy2(SKILL/'LICENSE.txt',target/'LICENSE.txt')
    (target/'selection.json').write_text(json.dumps(selected,indent=2)+'\n',encoding='utf-8')
    (target/'ATTRIBUTION.txt').write_text('Dazzler original SVG assets. Copyright 2026 Jon Gosier. Apache-2.0.\nUse semantic labels for meaningful icons; hide decorative SVGs from assistive technology.\n',encoding='utf-8')
    return selected


def plan_change(root, relative, proposed, destination):
    original=inside(root,relative); replacement=Path(proposed)
    if not original.is_file() or not replacement.is_file():raise ValueError('Both versions must be existing files')
    # One-file changes make apply/revert atomic and avoid partial multi-file migrations.
    old=original.read_bytes();new=replacement.read_bytes()
    old_text=old.decode('utf-8');new_text=new.decode('utf-8')
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=False)
    (destination/'before.txt').write_bytes(old);(destination/'after.txt').write_bytes(new)
    record={'schemaVersion':1,'root':str(Path(root).resolve()),'path':relative,'before':hashlib.sha256(old).hexdigest(),'after':hashlib.sha256(new).hexdigest(),'applied':False}
    (destination/'change.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
    import difflib
    diff=difflib.HtmlDiff(wrapcolumn=80).make_file(old_text.splitlines(),new_text.splitlines(),fromdesc='Current',todesc='Proposed',context=True)
    diff=diff.replace('<title></title>','<title>Dazzler change preview</title>')
    (destination/'preview.html').write_text(diff,encoding='utf-8')
    return record


def apply_change(folder, revert=False):
    folder=Path(folder);record=read(folder/'change.json');target=inside(record['root'],record['path'])
    expected=record['after'] if revert else record['before'];data=folder/('before.txt' if revert else 'after.txt')
    wanted=record['before'] if revert else record['after']
    if digest(target)!=expected:raise ValueError('Project changed since preview/apply; refusing to overwrite edits')
    if digest(data)!=wanted:raise ValueError('Preview content changed; regenerate the plan')
    # Write next to target and replace atomically; failure does not erase original.
    import tempfile,os
    with tempfile.NamedTemporaryFile(dir=target.parent,delete=False) as out:
        out.write(data.read_bytes());temporary=Path(out.name)
    try:os.replace(temporary,target)
    finally:
        if temporary.exists():temporary.unlink()
    record['applied']=not revert
    (folder/'change.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
    return record


def export_document(content, system, destination, mode='document'):
    """Portable HTML editions with optional native DOCX/PPTX from installed libraries."""
    if mode not in ('document','slides','docx','pptx'):raise ValueError('Unknown document mode')
    title=str(content.get('title','Untitled'));sections=content.get('sections',[])
    if not isinstance(sections,list) or not sections:raise ValueError('Provide nonempty sections')
    for section in sections:
        if not isinstance(section,dict) or not isinstance(section.get('body',''),str):raise ValueError('Each section requires a string body')
    tokens=system['palette']['modes']['light']['tokens'];font=system['fonts']['body'];heading=system['fonts']['heading']
    if any(not re.fullmatch(r'#[0-9a-fA-F]{6}',tokens[k]) for k in ('text','background','border')):raise ValueError('Document colors must be opaque hex tokens')
    destination=Path(destination).resolve()
    if destination.is_relative_to(SKILL.resolve()):raise ValueError('Export outside the installed skill')
    if destination.exists():raise ValueError('Refusing to overwrite an existing deliverable')
    embedded=False
    if mode=='docx':
        try:
            from docx import Document
            from docx.shared import Inches,Pt,RGBColor
        except ImportError as exc:raise RuntimeError('Native DOCX needs python-docx in the host runtime. Use document HTML when unavailable.') from exc
        doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.8)
        normal=doc.styles['Normal'];normal.font.name=font;normal.font.size=Pt(11)
        normal.font.color.rgb=RGBColor.from_string(tokens['text'].lstrip('#'))
        for level in ['Title','Heading 1','Heading 2']:
            doc.styles[level].font.name=heading;doc.styles[level].font.color.rgb=RGBColor.from_string(tokens['text'].lstrip('#'))
        doc.add_heading(title,0)
        for section in sections:
            if section.get('pageBreak'):doc.add_page_break()
            doc.add_heading(str(section.get('heading','')),1)
            for paragraph in section.get('body','').split('\n\n'):doc.add_paragraph(paragraph)
            if section.get('table'):
                rows=section['table'];table=doc.add_table(rows=0,cols=max(map(len,rows)));table.style='Light Shading Accent 1'
                for row in rows:
                    cells=table.add_row().cells
                    for i,value in enumerate(row):cells[i].text=str(value)
        # Apply brand role colors directly; do not inherit Word's default blue accent.
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        for table in doc.tables:
            for ri,row in enumerate(table.rows):
                for cell in row.cells:
                    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),tokens['background'][1:]);cell._tc.get_or_add_tcPr().append(shd)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:run.font.color.rgb=RGBColor.from_string(tokens['text'][1:]);run.bold=ri==0
        doc.save(destination)
    elif mode=='pptx':
        try:
            from pptx import Presentation
            from pptx.util import Inches,Pt
            from pptx.dml.color import RGBColor
        except ImportError as exc:raise RuntimeError('Native PPTX needs python-pptx in the host runtime. Use slides HTML when unavailable.') from exc
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        for section in [{'heading':title,'body':content.get('subtitle','')},*sections]:
            # Split long paragraphs into continuation slides, avoiding silent clipping.
            words=section.get('body','').split();chunks=[' '.join(words[i:i+65]) for i in range(0,len(words),65)] or ['']
            for i,body in enumerate(chunks):
                slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(tokens['background'][1:])
                for x,y,w,h,text,size,family in [(0.7,.6,11.9,1.2,str(section.get('heading',''))+(' (continued)' if i else ''),30,heading),(.7,2,11.9,4.5,body,24,font)]:
                    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=box.text_frame;frame.word_wrap=True
                    frame.text=text
                    for p in frame.paragraphs:p.font.name=family;p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(tokens['text'][1:])
            if section.get('table'):
                rows=section['table'];columns=max(map(len,rows))
                if columns>6:raise ValueError('Split tables wider than six columns into separate slide tables')
                for offset in range(1,len(rows),6):
                    page_rows=[rows[0],*rows[offset:offset+6]]
                    slide=prs.slides.add_slide(prs.slide_layouts[6])
                    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(tokens['background'][1:])
                    title_box=slide.shapes.add_textbox(Inches(.7),Inches(.5),Inches(12),Inches(.7));title_box.text=str(section.get('heading',''))+' — table'
                    title_box.text_frame.paragraphs[0].font.size=Pt(26)
                    title_box.text_frame.paragraphs[0].font.name=heading
                    title_box.text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string(tokens['text'][1:])
                    table=slide.shapes.add_table(len(page_rows),columns,Inches(.7),Inches(1.6),Inches(12),Inches(4.8)).table
                    for ri,row in enumerate(page_rows):
                        for ci,value in enumerate(row):
                            cell=table.cell(ri,ci);cell.text=str(value)
                            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(tokens['background'][1:])
                            for p in cell.text_frame.paragraphs:p.font.name=font;p.font.size=Pt(16);p.font.color.rgb=RGBColor.from_string(tokens['text'][1:]);p.font.bold=ri==0
        prs.save(destination)
    else:
        def esc(x):return html.escape(str(x),quote=True)
        # Font names are CSS string literals; prevent closing the style element.
        def css_font(x):return json.dumps(x).replace('<','\\3c ')
        slides=mode=='slides';pages=[]
        html_sections=[]
        for section in sections:
            if slides:
                words=section.get('body','').split();chunks=[' '.join(words[i:i+65]) for i in range(0,len(words),65)] or ['']
                for i,chunk in enumerate(chunks):html_sections.append({'heading':str(section.get('heading',''))+(' (continued)' if i else ''),'body':chunk})
                if section.get('table'):
                    rows=section['table']
                    for i in range(1,len(rows),6):html_sections.append({'heading':str(section.get('heading',''))+' — table','table':[rows[0],*rows[i:i+6]]})
            else:html_sections.append(section)
        for section in html_sections:
            table=''
            if section.get('table'):
                rows=section['table'];table='<table>'+''.join('<tr>'+''.join(f'<{"th" if i==0 else "td"}>{esc(v)}</{"th" if i==0 else "td"}>' for v in row)+'</tr>' for i,row in enumerate(rows))+'</table>'
            pages.append(f'<section><h2>{esc(section.get("heading",""))}</h2>'+''.join(f'<p>{esc(p)}</p>' for p in section.get('body','').split('\n\n'))+table+'</section>')
        css='body{font-family:'+css_font(font)+',sans-serif;color:'+tokens['text']+';background:'+tokens['background']+';line-height:1.65;margin:0}h1,h2{font-family:'+css_font(heading)+',serif;line-height:1.2}'
        css+='main{max-width:900px;margin:auto;padding:48px}h1{font-size:3rem}h2{font-size:2rem}section{padding:24px 0}table{border-collapse:collapse;width:100%}th,td{text-align:left;padding:12px;border-bottom:1px solid '+tokens['border']+'} @page{size:A4;margin:20mm}@media print{main{padding:0}h2{break-after:avoid}tr{break-inside:avoid}}@media(max-width:600px){main{padding:24px}h1{font-size:2.3rem}}'
        if slides:css+='section{min-height:65vh;border-bottom:2px solid;scroll-margin:20px} @media print{@page{size:landscape}section{min-height:0;break-before:page;break-inside:avoid}}'
        notices=[]
        for directory in content.get('fontExports',[]):
            import base64
            from urllib.parse import unquote
            directory=Path(directory).resolve();selection=read(directory/'selection.json')
            catalog=read(SKILL/'references/font-catalog.json');family=next(f for f in catalog['fonts'] if f['id']==selection['family'])
            for file in family['support_files']:
                source=directory/Path(file['path']).name
                if not source.is_file() or digest(source)!=file['sha256']:raise ValueError('Missing or modified font notice')
                if source.suffix.lower() in ('.txt','.md'):notices.append(esc(source.read_text(encoding='utf-8')))
            font_css=(directory/'fonts.css').read_text(encoding='utf-8')
            if '</style' in font_css.lower():raise ValueError('Unsafe exported font CSS')
            def embed(match):
                source=inside(directory,unquote(match.group(1)))
                face=next((f for f in family['files'] if f['filename']==source.name),None)
                if not face or digest(source)!=face['sha256']:raise ValueError('Unverified exported font binary')
                mime='font/woff2' if source.suffix=='.woff2' else 'font/woff' if source.suffix=='.woff' else 'font/otf' if source.suffix=='.otf' else 'font/ttf'
                return 'url("data:'+mime+';base64,'+base64.b64encode(source.read_bytes()).decode()+'")'
            embedded_css=re.sub(r'url\("\./([^\"]+)"\)',embed,font_css)
            if re.search(r'@import|url\(\s*(?!"data:)',embedded_css,re.I):raise ValueError('Font CSS must contain only embedded local resources')
            css=embedded_css+'\n'+css;embedded=True
        appendix='<details><summary>Font licenses and sources</summary><pre style="white-space:pre-wrap">'+'\n\n'.join(notices)+'</pre></details>' if notices else ''
        destination.write_text(f'<!doctype html><html lang="{esc(content.get("lang","en"))}"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{esc(title)}</title><style>{css}</style><main><h1>{esc(title)}</h1>{"".join(pages)}{appendix}</main></html>',encoding='utf-8')
    return {'file':str(destination),'format':mode,'fontsEmbedded':embedded,'verification':'Exported, not visually certified. Check page/slide layout in the target renderer. Native formats reference fonts; HTML optionally embeds verified exported font files and notices.'}


def main():
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('brand');p.add_argument('source');p.add_argument('--out',required=True)
    p=sub.add_parser('assets');p.add_argument('--kind',choices=['icon','illustration'],required=True);p.add_argument('--mood',default='');p.add_argument('--ids',nargs='*');p.add_argument('--out',required=True)
    p=sub.add_parser('plan');p.add_argument('--root',required=True);p.add_argument('--file',required=True);p.add_argument('--proposed',required=True);p.add_argument('--out',required=True)
    for action in ('apply','revert'):p=sub.add_parser(action);p.add_argument('plan')
    p=sub.add_parser('export');p.add_argument('--content',required=True);p.add_argument('--system',required=True);p.add_argument('--format',choices=['document','slides','docx','pptx'],required=True);p.add_argument('--out',required=True)
    args=parser.parse_args()
    if args.command=='brand':
        result=brand(args.source)
        with open(args.out,'x',encoding='utf-8') as f:json.dump(result,f,indent=2)
    elif args.command=='assets':result=assets(args.kind,args.mood,args.out,args.ids)
    elif args.command=='plan':result=plan_change(args.root,args.file,args.proposed,args.out)
    elif args.command in ('apply','revert'):result=apply_change(args.plan,args.command=='revert')
    else:result=export_document(read(args.content),read(args.system),args.out,args.format)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    try:main()
    except (ValueError,RuntimeError,OSError) as exc:raise SystemExit(str(exc))
