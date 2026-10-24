import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

spec = importlib.util.spec_from_file_location(
    "native_office",
    Path(__file__).resolve().parents[1]
    / "skills/dazzler-frontend/scripts/native_office.py",
)
native = importlib.util.module_from_spec(spec)
spec.loader.exec_module(native)


class NativeTests(unittest.TestCase):
    def test_bad_data_never_becomes_chart_values(self):
        for values in [[True], [float("nan")], ["3"], []]:
            with self.assertRaises(ValueError):
                native.validate_chart(
                    {
                        "categories": ["A"],
                        "series": [{"name": "Count", "values": values}],
                    }
                )
        native.validate_chart(
            {"categories": ["A"], "series": [{"name": "Count", "values": [None]}]}
        )

    def test_structure_does_not_claim_rendered_fonts(self):
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / "a.docx"
            with zipfile.ZipFile(p, "w") as z:
                z.writestr("word/document.xml", "<document><t>Example</t></document>")
            result = native.inspect(p)
            self.assertEqual(result["fontVerification"], "unknown-until-rendered")
            self.assertEqual(result["textRuns"], 1)

    def test_native_chart_preserves_missing_and_can_update_workbook(self):
        try:
            from pptx import Presentation
            from pptx.chart.data import CategoryChartData
        except ImportError:
            self.skipTest("Optional python-pptx runtime unavailable")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.pptx"
            out = root / "chart.pptx"
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(source)
            original = source.read_bytes()
            data = {
                "categories": ["Jan", "Feb", "Mar"],
                "series": [{"name": "Observed", "values": [2, None, 5]}],
            }
            native.add_chart(source, out, data, 1, [1, 1, 5, 3], "line")
            self.assertEqual(source.read_bytes(), original)
            deck = Presentation(out)
            chart = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
            self.assertEqual(list(chart.series[0].values), [2, None, 5])
            self.assertEqual(
                chart._chartSpace.xpath(".//c:dispBlanksAs")[0].get("val"), "gap"
            )
            update = CategoryChartData()
            update.categories = ["Jan", "Feb", "Mar"]
            update.add_series("Observed", [8, None, 5])
            chart.replace_data(update)
            edited = root / "edited.pptx"
            deck.save(edited)
            changed = next(
                s.chart for s in Presentation(edited).slides[0].shapes if s.has_chart
            )
            self.assertEqual(list(changed.series[0].values), [8, None, 5])
            with zipfile.ZipFile(edited) as z:
                self.assertTrue(
                    any(n.startswith("ppt/embeddings/") for n in z.namelist())
                )
            self.assertEqual(native.inspect(edited)["nativeCharts"], 1)
            with self.assertRaises(FileExistsError):
                native.add_chart(source, out, data, 1, [1, 1, 5, 3], "line")


if __name__ == "__main__":
    unittest.main()
