"""Context-specific, fictional examples for editable document templates."""
import html
from datetime import datetime, timezone

def p(text):return {'type':'p','text':text}
def h(text):return {'type':'h','text':text}
def call(label,text):return {'type':'callout','label':label,'text':text}
def tab(headers,rows,widths=None):return {'type':'table','headers':headers,'rows':rows,'widths':widths}
def bullets(*items):return {'type':'list','items':items}
def page(title,*blocks,kicker=None):return {'title':title,'blocks':blocks,'kicker':kicker}

DOCS=[
 dict(id='professional',title='Customer portal launch',subtitle='Weekly delivery report',accent='#244D70',tint='#EAF1F7',font='Arial',tag='DELIVERY OFFICE',use='Weekly executive status report with decision, milestones, risks and ownership',pages=[page('Customer portal launch',
 p('Reporting period  14–18 June 2027     •     Prepared by Maya Chen, Delivery Lead'),
 call('AMBER — decision needed by 21 June','The 12 July launch remains achievable. Identity integration is three days behind; approving a limited pilot keeps the public launch date intact.'),
 tab(['Scope complete','Budget used','Next milestone'],[['18 of 24 deliverables','USD 42,000 of 60,000','Pilot acceptance • 25 June']]),
 h('Delivery against the plan'),tab(['Workstream','Owner','Status and evidence'],[['Account setup','Alex Rivera','Complete • 24 acceptance checks passed'],['Invoice history','Sam Okafor','In review • finance sign-off due 21 June'],['Single sign-on','Maya Chen','At risk • sandbox credentials still pending']], [25,20,55]),
 h('Risks and response'),p('R-03 • Identity provider access. If credentials arrive after 22 June, the pilot will use password sign-in. Maya owns the fallback; security review is booked for 23 June.'),
 h('Decision for the sponsor'),p('Approve the 20-customer pilot with password sign-in, subject to security review. No additional budget is requested. Priya Shah to confirm by 21 June, 15:00.'),
 h('Next week'),bullets('Finance: approve invoice totals and download wording by 21 June.','Engineering: complete pilot deployment and rollback rehearsal by 24 June.','Customer success: confirm pilot participants and support coverage by 25 June.'))]),
 dict(id='legal',title='Supplier exit review',subtitle='Matter assessment memorandum',accent='#393B40',tint='#F0F0EE',font='Georgia',tag='MATTER 27-014',use='Legal matter memo with chronology, evidence gaps and an authority research matrix',pages=[page('Supplier exit review',
 tab(['To','From','Date'],[['Reviewing counsel','Jordan Lee • Legal operations','18 June 2027']]),
 p('Client matter: Alder Studio / Northline Services. Purpose: organize the record for counsel’s review of a proposed supplier transition.'),
 call('QUESTION FOR REVIEW','What contractual notice, transition and data-return requirements must be verified before the client sets an exit date?'),
 h('Preliminary position'),p('The file is incomplete. The team should obtain the executed agreement, amendments and delivery record before recording a notice deadline or sending an exit communication. This example does not reach a legal conclusion.'),
 h('Working chronology'),tab(['Date','Recorded event','Evidence'],[['3 February','Services began','Onboarding email, file E-01'],['7 May','Client reported delayed handover','Support ticket, file E-02'],['15 June','Operations requested exit options','Internal instruction, file E-03']], [19,46,35]),
 h('Facts still to establish'),bullets('Which signed terms govern, and have they been amended?','Who is authorized to give and receive notice, and by which method?','What client data or work product is held by the supplier?'),
 p('Evidence status: the chronology above is fictional sample material. Replace each reference with the actual retained source.')),page('Analysis and review plan',
 h('Contract and authority matrix'),tab(['Issue','Source to verify','Analysis to complete'],[['Notice','Executed agreement and amendments','Trigger, period, recipient and delivery method'],['Transition','Exit and assistance provisions','Services, duration, fees and dependencies'],['Data return','Data processing terms','Format, retention and deletion obligations'],['Governing law','Agreement and current authorities','Insert verified jurisdiction-specific research']], [20,37,43]),
 h('Competing considerations'),p('Operations wants a short transition; the technology team needs time to validate exports. Counsel should assess the verified terms against both needs and identify any required approvals or unresolved questions.'),
 h('Requested evidence'),tab(['Item','Owner','Target'],[['Executed agreement and amendments','Procurement','21 June'],['Service and notice correspondence','Operations','21 June'],['Data inventory and export test','Technology','23 June']]),
 h('Review and disposition'),p('Reviewing counsel: [Name]     Review date: [Date]\nApproved next action: [Insert after review]\nAuthorities and pinpoint references: [Add verified sources]'),
 p('Layout example only. No jurisdiction-specific authorities, operative clauses or privilege designation are supplied.'))]),
 dict(id='business',title='Customer onboarding redesign',subtitle='Proposal for Alder Studio',accent='#185849',tint='#EAF3EE',font='Arial',tag='NORTHSTAR DESIGN • PROPOSAL 027',use='Two-page services proposal with deliverables, exclusions, pricing and acceptance',pages=[page('Customer onboarding redesign',
 p('Prepared for Alder Studio     •     Prepared by Northstar Design\nIssued 18 June 2027     •     Proposal valid through 2 July 2027'),
 call('THE OUTCOME','A clearer first-week experience: one welcome journey, one account checklist and fewer handoffs between sales and customer success.'),
 h('What we heard'),p('New customers receive setup instructions in three separate emails. The proposed work consolidates those steps, clarifies ownership and gives the service team a reusable launch checklist.'),
 h('Scope and delivery'),tab(['Phase','Deliverables','Schedule'],[['Discover','Five staff interviews; journey map; agreed problem statement','Week 1'],['Design','Welcome flow; six screen designs; email copy outline','Weeks 2–3'],['Handover','Annotated prototype; component notes; two training sessions','Week 4']], [18,61,21]),
 h('How we will work'),bullets('One 30-minute review each Tuesday, with a named client decision-maker.','Two consolidated feedback rounds on the prototype.','Shared decision log and weekly written delivery update.'),
 h('Outside this scope'),p('Production code, system migration, translation and new brand identity work are excluded. Any added scope requires a written change agreement.')),page('Investment and approval',
 tab(['Work package','Fee in USD'],[['Research and journey mapping','$3,000'],['Experience design and prototype','$7,500'],['Handover and training','$1,500'],['Total project fee','$12,000']], [72,28]),
 h('Payment schedule'),p('40% at kickoff ($4,800), 40% at prototype review ($4,800), and 20% at handover ($2,400). This illustrative price excludes applicable taxes and third-party expenses.'),
 h('Client responsibilities'),bullets('Provide current onboarding material and access to five interview participants.','Return consolidated feedback within two working days.','Confirm content accuracy and nominate an approver before kickoff.'),
 h('Acceptance criteria'),p('Handover is complete when the agreed six screens, email outline, annotated prototype and training sessions have been delivered. The project lead records acceptance or a specific list of scope-related corrections.'),
 h('Proposed next step'),p('Confirm the scope and proposed 5 July kickoff. The parties should finalize their service agreement before work begins.'),
 tab(['Client approver','Provider approver'],[['[Name and role]\n[Signature and date]','[Name and role]\n[Signature and date]']]),
 p('Sample proposal. Replace the fictional parties, dates, pricing and terms with approved project details.'))]),
 dict(id='fun',title='The great game night',subtitle='Pick a team • Bring your best terrible drawing',accent='#80396A',tint='#F9EBF3',font='Arial',tag='YOU ARE INVITED',use='A playful event invitation with actual logistics, schedule and RSVP details',center=True,pages=[page('The great game night',
 call('SATURDAY 24 JULY','6:30–10:00 PM  •  The Rivera living room\n18 Willow Lane, Sampletown — fictional venue'),
 p('Come for the pizza. Stay for the dramatic final round. We are mixing quick team games with a few old favorites, and there is a seat for complete beginners.'),
 tab(['Time','What is happening'],[['6:30','Arrive, choose a team and grab a slice'],['7:00','Drawing game warm-up'],['7:45','Team trivia and snack break'],['8:45','Choose-your-own table games'],['9:45','One last round and the extremely unofficial awards']], [20,80]),
 h('Bring one thing'),p('A favorite game, a snack to share, or just yourself. Comfortable clothes encouraged. Please tell us about food needs when you reply.'),
 h('Save your seat'),p('RSVP to Alex by 20 July at games@example.com. Let us know how many people are coming so we can plan teams and dinner.'),
 p('All ages welcome with a grown-up. Step-free entry through the side gate. Quiet room available away from the games.'),
 call('HOUSE RULE','Be kind, explain the rules, and celebrate a good move—even when it is not yours.'))]),
 dict(id='family',title='Our week at a glance',subtitle='Morgan household • 21–27 June 2027',accent='#725031',tint='#F7EFE4',font='Arial',tag='THE FRIDGE PLANNER',landscape=True,use='Seven-day household schedule with meals, pickups, responsibilities and shopping',pages=[page('Our week at a glance',
 tab(['Day','Plans and appointments','Pickup or owner','Dinner'],[['Mon 21','Library returns • 16:00','Alex picks up Riley','Vegetable pasta'],['Tue 22','Swimming • 17:15','Taylor drives; Alex collects','Rice bowls'],['Wed 23','School project due','Riley packs model Tuesday','Soup and toast'],['Thu 24','Dentist • 15:30','Taylor • leave at 15:00','Tray-bake potatoes'],['Fri 25','Movie night • 19:00','Everyone chooses one option','Homemade pizza'],['Sat 26','Park picnic • 12:00','Alex packs lunch','Leftovers'],['Sun 27','Plan next week • 17:30','All • 15-minute check-in','Family choice']], [12,38,28,22]),
 h('Shared jobs'),tab(['Alex','Taylor','Riley'],[['Shopping • Tuesday\nRecycling • Thursday','Laundry • Wednesday\nCheck calendar • Sunday','Water plants • Mon/Wed/Fri\nPack bag each evening']]),
 h('Shopping and reminders'),p('Milk • oats • tomatoes • apples • rice • dishwasher tablets\nReturn library books Monday. Swimming bag by the door Tuesday. Confirm the picnic weather Saturday morning.'),
 call('ROOM FOR CHANGE','If a plan changes, tell the person doing pickup first. Add a new plan here: [Family note]'))]),
 dict(id='presentation',title='Approve the onboarding pilot',subtitle='Decision brief • Leadership review • 18 June 2027',accent='#353F87',tint='#EDEFFC',font='Arial',tag='ALDER STUDIO',landscape=True,use='Three-page landscape decision presentation with evidence, options and an action plan',pages=[page('Approve the onboarding pilot',
 call('THE ASK','Approve a four-week pilot for 20 new customers, with a USD 6,000 spend limit and a named customer-success owner.'),
 tab(['Observed friction','Proposed change','Decision today'],[['Setup instructions span three emails','One checklist and one welcome sequence','Authorize the pilot and owner']], [33,34,33]),
 h('Why this is worth testing'),p('The pilot will test whether customers can finish setup with fewer support handoffs. It is an experiment with a defined end date, not a commitment to a full rollout.'),
 p('Presenter note: open with the decision, then explain what the team will learn. All figures and organizations in this brief are illustrative.')),page('Compare the options',
 tab(['Option','Benefit','Trade-off'],[['Keep the current process','No near-term delivery cost','Existing handoffs remain'],['Pilot with 20 customers','Tests the approach before a broad change','Requires four weeks and a dedicated owner'],['Rebuild immediately','Moves directly to a single experience','Commits more resources before learning']], [24,38,38]),
 h('Recommended option'),p('Run the limited pilot. Use the current system, rewrite the welcome sequence and add one checklist; defer any platform migration.'),
 tab(['Measure','Pilot target','Source'],[['Setup completed within 7 days','At least 16 of 20 participants','Checklist event log'],['Support contacts in first week','Record per participant; compare with baseline','Support ticket report'],['Serious account issues','Zero unresolved at closeout','Incident log']]),
 p('Presenter note: targets are proposed thresholds for this sample, not reported results.')),page('Make the next month measurable',
 tab(['Week','Work','Accountable owner'],[['1','Finalize copy and baseline measures','Maya • Customer success'],['2','Invite participants and start the pilot','Sam • Operations'],['3','Review exceptions and fix copy gaps','Alex • Product'],['4','Report findings and recommend next step','Maya • Customer success']]),
 h('Guardrails'),bullets('Pause new invitations if account access is affected.','Keep the existing onboarding path available during the test.','Do not expand the pilot without reviewing the final evidence.'),
 call('RECORD THE DECISION','Sponsor: [Name]    Outcome: [Approve / revise / decline]\nOwner confirmed: [Name]    Review meeting: [Date]'))]),
 dict(id='school',title='How light affects seedling growth',subtitle='Science investigation • Year 8',accent='#285576',tint='#EAF3FA',font='Arial',tag='STUDENT RESEARCH REPORT',use='Science report with variables, repeatable method, sample data and reflection',pages=[page('How light affects seedling growth',
 p('Student: [Name]    Class: [Class]    Teacher: [Name]    Date: [Date]'),
 h('Question and prediction'),p('How does daily light exposure affect bean seedling height over two weeks? I predict that the group receiving more light will show healthier leaf growth. Height alone may not capture plant health.'),
 h('Variables'),tab(['Type','What is measured or controlled'],[['Independent','Daily light exposure: 4 hours or 8 hours'],['Dependent','Seedling height in centimeters; leaf appearance'],['Controlled','Seed type, soil, container, water and measurement time']], [24,76]),
 h('Materials and method'),bullets('Prepare six similar seedlings in identical containers; label three A and three B.','Place group A in 4 hours of light and group B in 8 hours. Keep other conditions consistent.','Water each container with the same measured amount daily.','Measure height from soil level at the same time on days 0, 7 and 14. Record observations.'),
 h('Care and safety'),p('Ask a teacher before using lamps. Keep water away from electrical equipment. Wash hands after handling soil.'),
 call('RECORD HONESTLY','The next page contains example data to demonstrate the layout. Replace it with your own measurements; do not present it as an experiment you performed.')),page('Results and reflection',
 tab(['Group','Day 0 mean','Day 7 mean','Day 14 mean'],[['A • 4 hours','3.0 cm','5.4 cm','8.1 cm'],['B • 8 hours','3.1 cm','5.8 cm','9.2 cm']]),
 p('Illustrative means for three plants per group. Retain individual measurements in your notebook so another reader can check the calculation.'),
 h('What the example suggests'),p('In this sample, group B grew 6.1 cm and group A grew 5.1 cm. The difference is 1.0 cm. With only three plants per group, this does not establish a general rule.'),
 h('Limitations and improvements'),bullets('Record leaf color and leaf count as well as height.','Rotate container positions to reduce uneven light exposure.','Repeat with more seedlings and record individual variation.'),
 h('Sources and next question'),p('Sources consulted: [Author, title, publication date and page or URL]\nNext investigation: would the same pattern appear with another plant species?'),
 h('Teacher feedback'),p('[Strengths]\n[One specific improvement]'))]),
 dict(id='marketing',title='Make the first visit easy',subtitle='Campaign brief • Juniper neighborhood launch',accent='#993E28',tint='#FFF0E9',font='Arial',tag='CREATIVE BRIEF 04',use='Campaign brief with audience, message, channel budget, production details and measurement',pages=[page('Make the first visit easy',
 p('Owner: Lena Ortiz    Campaign: 1–31 July 2027    Review: 18 June 2027'),
 call('PRIMARY OBJECTIVE','Generate 120 qualified reservation inquiries during the launch month. This is a sample target, not a forecast or an achieved result.'),
 h('Who we are speaking to'),p('People living or working within the local neighborhood who want a relaxed place for a midweek dinner. They need to understand the menu, typical spend and how to book before deciding.'),
 h('Message and response'),tab(['Main message','Supporting content','Call to action'],[['A seasonal dinner close to home','Menu highlights, room details, service hours and booking instructions','View the menu and request a table']]),
 h('Creative direction'),p('Warm food-led photography, short ingredient descriptions and readable menu typography. Show the actual space when photographs are supplied. Avoid invented reviews, awards or claims about local sourcing.'),
 h('Required information'),bullets('Approved menu and prices; service dates and hours.','A clear booking route and accurate location information.','Confirmed dietary and accessibility guidance from the operator.'),
 h('Approval route'),p('Chef reviews menu descriptions. Operations confirms hours and booking information. Campaign owner approves final copy before publication.')),page('Channels and measurement',
 tab(['Channel','Asset and purpose','Budget USD'],[['Local social','Three short edits leading to the menu','$1,200'],['Search','Intent-led booking queries','$600'],['Print','Neighborhood cards with a trackable link','$300'],['Production','Photography and copy adaptation','$900'],['Total','Illustrative allocation','$3,000']], [23,56,21]),
 h('Production schedule'),tab(['Date','Delivery','Approver'],[['22 June','Menu and message approved','Chef + operations'],['25 June','Photography and copy ready','Campaign owner'],['28 June','Tracking and booking checks','Web owner'],['1 July','Campaign begins','Campaign owner']]),
 h('Measurement plan'),p('Track menu visits, reservation inquiry starts and completed inquiries separately. Use distinct campaign links. Report weekly inquiry counts and spend; do not treat a click as a booking.'),
 h('Decision after week one'),p('If people reach the menu but abandon the inquiry form, inspect the form before increasing spend. Record changes so performance can be interpreted against the correct version.'),
 p('Replace the fictional restaurant, budget, targets and dates with approved campaign details.'))]),
 dict(id='restaurant',title='Juniper dinner menu',subtitle='Early summer • Dinner from 17:30',accent='#4B5633',tint='#F5F2E5',font='Georgia',tag='JUNIPER • SEASONAL KITCHEN',center=True,use='Restaurant menu with courses, ingredient descriptions, prices and service notes',pages=[page('Juniper dinner menu',
 p('An unhurried dinner built around vegetables, grains and the season.\nIllustrative menu • Prices in USD'),
 h('To begin'),tab(['Dish','Preparation','Price'],[['Garden leaves','Radish, soft herbs, mustard dressing','14'],['Roasted beetroot','Whipped ricotta, orange, toasted hazelnuts','16'],['Warm sourdough','Cultured butter and sea salt','8']], [29,58,13]),
 h('From the kitchen'),tab(['Dish','Preparation','Price'],[['Summer squash','White beans, sage, pumpkin seeds','26'],['Roast chicken','New potatoes, greens, pan jus','32'],['Market fish','Braised fennel, lemon, herb oil','34']], [29,58,13]),
 h('Something sweet'),tab(['Dish','Preparation','Price'],[['Pear and almond','Warm pear, almond crumb, cream','11'],['Dark chocolate','Chocolate crémeux, olive oil, sea salt','12']], [29,58,13]),
 h('At your table'),p('Please speak with the team about allergies before ordering. Ingredients and preparation may change. Add the restaurant’s verified allergen, service-charge and tax information before using this menu.'),
 p('Reservations: hello@example.com    •    Wednesday–Sunday, 17:30–22:00\nSample restaurant details for template demonstration.'))]),
 dict(id='technical',title='Order status webhook delivery',subtitle='Technical design • RFC 014 • Proposed',accent='#305666',tint='#EAF3F5',font='Arial',tag='PLATFORM ENGINEERING',use='Technical RFC with delivery guarantees, payload, failure handling, tests and rollout',pages=[page('Order status webhook delivery',
 p('Owner: Alex Rivera    Reviewers: API + Reliability    Updated: 18 June 2027'),
 call('DESIGN DECISION','Deliver order status events asynchronously with at-least-once delivery. Consumers deduplicate on event_id; the service does not promise exactly-once processing.'),
 h('Context and boundaries'),p('Partners currently poll for order updates. This design adds signed event delivery to registered HTTPS endpoints. It excludes partner endpoint provisioning, historical backfills and delivery to arbitrary customer-supplied URLs.'),
 h('Delivery path'),p('Order transaction → transactional outbox → delivery queue → worker → registered partner endpoint. A delivery record stores each attempt and its result.'),
 tab(['Requirement','Design response'],[['No lost committed events','Write the outbox record in the order transaction'],['Duplicate-safe processing','Keep event_id stable across retries'],['Endpoint authenticity','Sign the raw request body with the partner secret'],['Failure isolation','Separate per-partner concurrency and delivery limits']], [34,66]),
 h('Event example'),{'type':'code','text':'{\n  "event_id": "evt_demo_0042",\n  "type": "order.shipped",\n  "occurred_at": "2027-06-18T14:05:00Z",\n  "data": {"order_id": "ord_demo_108", "status": "shipped"}\n}'},
 p('The identifiers and payload are illustrative. No credentials or customer information are included.')),page('Failures and release plan',
 tab(['Response','Worker behavior'],[['2xx','Mark delivered; retain attempt metadata'],['408, 429 or 5xx','Retry with bounded exponential backoff and jitter'],['Other 4xx','Stop retrying; flag configuration review'],['Timeout after 10 seconds','Treat as unknown outcome; retry with the same event_id']], [31,69]),
 h('Retry and security boundaries'),p('Limit delivery to six attempts within 24 hours. Revalidate registered destinations and block private network targets. Rotate secrets through the existing credential service; never put secrets in event payloads or logs.'),
 h('Acceptance checks'),bullets('A committed order change creates one durable outbox record.','A worker restart preserves pending events and stable event IDs.','A duplicate event is harmless in the reference consumer.','A failing partner endpoint does not delay healthy partners.'),
 h('Rollout and rollback'),tab(['Stage','Gate'],[['Internal endpoint','Replay fixtures; test timeouts and invalid signatures'],['Two pilot partners','Review delivery latency and retry rate for 48 hours'],['General enablement','Expand gradually after operational sign-off'],['Rollback','Stop new dispatch; preserve queue and outbox for replay']], [31,69]),
 h('Operational ownership'),p('Reliability owns alerts for oldest pending event age and repeated partner failures. API owns the payload contract and versioning. Thresholds must be set from pilot measurements before general enablement.'))]),
]

def render_docx(d,path):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(11 if d.get('landscape') else 8.5);sec.page_height=Inches(8.5 if d.get('landscape') else 11)
    if d.get('landscape'):sec.orientation=WD_ORIENT.LANDSCAPE
    sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.72)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Normal',10),('Title',28),('Subtitle',12),('Heading 1',13),('Heading 2',11),('List Bullet',10)]:
        st=doc.styles[name];st.font.name=d['font'];st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts
        for key in list(rf.attrib):
            if 'Theme' in key:del rf.attrib[key]
        st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(6)
        if name=='Heading 1':st.paragraph_format.space_before=Pt(12)
    doc.core_properties.title=d['title'];doc.core_properties.author='Dazzler / Jon Gosier';doc.core_properties.subject=d['use'];doc.core_properties.comments='Fictional worked template. Replace sample details. Apache-2.0.'
    doc.core_properties.created=doc.core_properties.modified=datetime(2026,9,28,tzinfo=timezone.utc)
    footer=sec.footer.paragraphs[0];footer.text='Dazzler sample template • '+d['id']+' • Replace fictional details before use';footer.runs[0].font.size=Pt(8)
    for pi,pg in enumerate(d['pages']):
        if pi:doc.add_page_break()
        kicker=doc.add_paragraph(d['tag']+('  /  '+str(pi+1) if len(d['pages'])>1 else ''));kicker.runs[0].font.size=Pt(9);kicker.runs[0].bold=True;kicker.runs[0].font.color.rgb=RGBColor.from_string(d['accent'][1:])
        title=doc.add_paragraph(pg['title'],'Title')
        if d.get('center'):title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if pi==0:doc.add_paragraph(d['subtitle'],'Subtitle')
        for b in pg['blocks']:
            kind=b['type']
            if kind in ('p','h'):doc.add_paragraph(b['text'],'Heading 1' if kind=='h' else 'Normal')
            elif kind=='list':
                for text in b['items']:doc.add_paragraph(text,'List Bullet')
            elif kind=='code':
                para=doc.add_paragraph(b['text']);para.runs[0].font.name='Consolas';para.runs[0].font.size=Pt(9)
            elif kind=='callout':
                para=doc.add_paragraph();para.paragraph_format.space_before=Pt(8);para.paragraph_format.space_after=Pt(10)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),d['tint'][1:]);para._p.get_or_add_pPr().append(shade)
                para.add_run(b['label']+'\n').bold=True;para.add_run(b['text'])
            elif kind=='table':
                table=doc.add_table(rows=0,cols=len(b['headers']));table.autofit=False
                fractions=b.get('widths') or [100/len(b['headers'])]*len(b['headers']);width=sec.page_width-sec.left_margin-sec.right_margin
                for col,frac in zip(table.columns,fractions):col.width=int(width*frac/100)
                for ri,row in enumerate([b['headers'],*b['rows']]):
                    cells=table.add_row().cells
                    no_split=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                    if ri==0:table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    for ci,(cell,text) in enumerate(zip(cells,row)):
                        cell.width=int(width*fractions[ci]/100);cell.text=str(text);props=cell._tc.get_or_add_tcPr()
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),d['tint'][1:] if ri%2==0 else 'FFFFFF');props.append(shade)
                        margins=OxmlElement('w:tcMar')
                        for side in ['top','left','bottom','right']:
                            val=OxmlElement('w:'+side);val.set(qn('w:w'),'80');val.set(qn('w:type'),'dxa');margins.append(val)
                        props.append(margins)
                        for para in cell.paragraphs:
                            para.paragraph_format.space_after=Pt(1)
                            for run in para.runs:run.bold=ri==0;run.font.size=Pt(9)
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.save(path)

def render_html(d,path):
    e=lambda x:html.escape(str(x),quote=True)
    def block(b):
        k=b['type']
        if k in ('p','h'):return f'<{"h2" if k=="h" else "p"}>{e(b["text"])}</{"h2" if k=="h" else "p"}>'
        if k=='list':return '<ul>'+''.join('<li>'+e(x)+'</li>' for x in b['items'])+'</ul>'
        if k=='code':return '<pre><code>'+e(b['text'])+'</code></pre>'
        if k=='callout':return '<aside class="callout"><strong>'+e(b['label'])+'</strong><p>'+e(b['text'])+'</p></aside>'
        if d['id']=='restaurant':return '<div class="menu-course">'+''.join('<div class="dish"><strong>'+e(n)+'<span>'+e(price)+'</span></strong><p>'+e(desc)+'</p></div>' for n,desc,price in b['rows'])+'</div>'
        return '<div class="table-wrap"><table><thead><tr>'+''.join('<th scope="col">'+e(x)+'</th>' for x in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+e(v)+'</td>' for v in row)+'</tr>' for row in b['rows'])+'</tbody></table></div>'
    pages=''.join('<article class="sheet"><p class="kicker">'+e(d['tag'])+' / '+str(i+1)+'</p><h1>'+e(pg['title'])+'</h1>'+('<p class="subtitle">'+e(d['subtitle'])+'</p>' if i==0 else '')+''.join(block(b) for b in pg['blocks'])+'<footer>Illustrative Dazzler template • Replace fictional details before use</footer></article>' for i,pg in enumerate(d['pages']))
    font='Young Serif' if d['font']=='Georgia' else 'Work Sans'
    css=f":root{{--accent:{d['accent']};--tint:{d['tint']};--title:'{font}';}}"+'''*{box-sizing:border-box}body{margin:0;background:var(--tint);color:#20252C;font:16px/1.55 'Work Sans',Arial,sans-serif}a{color:inherit}button{font:inherit;background:var(--accent);color:white;border:0;padding:10px 18px;cursor:pointer;border-radius:4px}:focus-visible{outline:3px solid #20252C;outline-offset:4px}.tools{max-width:960px;margin:24px auto;display:flex;gap:20px;align-items:center}.sheet{max-width:960px;margin:24px auto;background:white;padding:54px 62px;box-shadow:0 8px 30px #20252C0A}h1{font:500 37px/1.15 var(--title),Georgia,serif;letter-spacing:-.035em;margin:15px 0}h2{font-size:19px;margin:25px 0 10px}p,td{white-space:pre-line}p{margin:12px 0}.subtitle{font-size:20px;color:#515C69}.kicker{font-size:11px;font-weight:700;letter-spacing:.1em;color:var(--accent)}.callout{background:var(--tint);padding:18px 22px;margin:24px 0}.callout strong{font-size:12px;letter-spacing:.04em}.callout p{margin:8px 0 0}table{border-collapse:collapse;width:100%;font-size:14px;margin:18px 0}th,td{padding:12px;text-align:left;border-bottom:1px solid #C8CDD0;vertical-align:top}th{background:var(--tint);font-size:12px}tr{break-inside:avoid}.table-wrap{overflow:auto}li{margin-bottom:8px}pre{font:13px/1.6 Consolas,monospace;background:#F1F3F5;padding:20px;overflow:auto}footer{font-size:11px;color:#515C69;margin-top:30px}@page{size:letter;margin:15mm}@media print{body{background:white;font-size:10pt}.tools{display:none}.sheet{box-shadow:none;margin:0;padding:0;max-width:none;break-after:page}.sheet:last-child{break-after:auto}h1{font-size:25pt}h2{font-size:13pt;break-after:avoid}table{font-size:9pt}th,td{padding:7px}.callout{padding:12px}p{orphans:3;widows:3}}@media(max-width:700px){.sheet{margin:0 0 20px;padding:28px 22px}.tools{padding:0 22px}h1{font-size:30px}td,th{padding:9px;font-size:12px}}'''
    if d.get('landscape'):css+='@page{size:letter landscape}.sheet{max-width:1180px}'
    if d.get('center'):css+='h1,.subtitle,.kicker{text-align:center}h1{font-size:48px}.sheet{max-width:800px}'
    if d['id']=='restaurant':css+='.dish{margin:20px 0}.dish strong{display:flex;justify-content:space-between;gap:24px}.dish p{font-size:14px;margin:3px 0;color:#515C69}h2{text-align:center;font:24px Young Serif,Georgia,serif;margin-top:34px}'
    path.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+e(d['title'])+' — Dazzler</title><link rel="stylesheet" href="../fonts/fonts.css"><style>'+css+'</style></head><body><div class="tools"><button onclick="window.print()">Print or save as PDF</button><span>Worked example • Editable source</span></div>'+pages+'</body></html>',encoding='utf-8')
