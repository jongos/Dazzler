"""Original same-content native Word proof for Dazzler's compositional decisions.
Not a production template catalog; each plan records independent choices.
"""

import argparse, colorsys, json, subprocess, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DATA = [18, 21, 24, 29, 34, 42]
TITLE = "Urban Canopy Program"
INTRO = "A six year review of shade coverage in a fictional city. This document compares the recorded coverage values and identifies the information needed before selecting new planting areas."
FINDING = "Coverage rose from 18% in 2020 to 42% in 2025, an increase of 24 percentage points. The final value is 2 points above the illustrative target of 40%."
LIMIT = "These values are synthetic design-test data. Coverage alone does not demonstrate equal access, tree survival or a causal health benefit. The target is illustrative, not an approved municipal commitment."
NEXT = "Before choosing planting locations, obtain neighborhood-level canopy coverage, heat exposure, land availability and maintenance capacity. Compare those measures together rather than ranking locations by coverage alone."


def shade(p, color):
    x = OxmlElement("w:shd")
    x.set(qn("w:fill"), color)
    p._p.get_or_add_pPr().append(x)


def border(p, color):
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "12")
    e.set(qn("w:color"), color)
    b.append(e)
    p._p.get_or_add_pPr().append(b)


def paragraph(doc, text, style=None, size=None, color=None, bold=False):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.keep_together = True
    for r in p.runs:
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r.bold = bold
    return p


def picture(doc, file, width, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    pic = p.add_run().add_picture(str(file), width=Inches(width))
    pic._inline.docPr.set("descr", label)
    return p


def illustration(out, d, accent, on_accent):
    im = Image.new("RGB", (1400, 390), "#" + accent)
    draw = ImageDraw.Draw(im)
    if d["motif"] == "data-derived":
        for i in range(100):
            x = 55 + (i % 25) * 51
            y = 65 + (i // 25) * 63
            draw.rounded_rectangle(
                (x, y, x + 33, y + 44),
                radius=16,
                fill="#" + on_accent if i < 42 else None,
                outline="#" + on_accent,
                width=2,
            )
    elif d["motif"] == "organic":
        for i in range(9):
            draw.arc(
                (i * 110 - 240, -160, i * 110 + 600, 630),
                195,
                345,
                fill="#" + on_accent,
                width=10,
            )
    elif d["motif"] == "geometry":
        for i in range(12):
            x = i * 135 - 50
            draw.polygon(
                [(x, 390), (x + 240, 0), (x + 320, 0), (x + 80, 390)],
                fill="#" + on_accent,
            )
    elif d["motif"] == "typographic":
        f = ImageFont.truetype(r"C:\Windows\Fonts\georgiab.ttf", 230)
        draw.text((55, 40), "18", font=f, fill="#" + on_accent)
        draw.text((730, 40), "42", font=f, fill="#" + on_accent)
        draw.line((430, 210, 670, 210), fill="#" + on_accent, width=12)
        draw.polygon([(670, 210), (625, 175), (625, 245)], fill="#" + on_accent)
    else:
        for i in range(16):
            draw.line(
                (70, 40 + i * 20, 400 + i * 57, 40 + i * 20),
                fill="#" + on_accent,
                width=4,
            )
    im.save(out)


def chart(out, accent, background, ink):
    im = Image.new("RGB", (1400, 430), "#" + background)
    dr = ImageDraw.Draw(im)
    f = ImageFont.truetype(r"C:\Windows\Fonts\calibri.ttf", 27)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\calibrib.ttf", 31)
    for y in [0, 20, 40]:
        py = 350 - y * 6
        dr.line((100, py, 1300, py), fill="#" + ink, width=2)
        dr.text((32, py - 15), str(y) + "%", font=f, fill="#" + ink)
    points = [(170 + i * 210, 350 - v * 6) for i, v in enumerate(DATA)]
    dr.line(points, fill="#" + accent, width=6)
    for i, (x, y) in enumerate(points):
        dr.ellipse((x - 7, y - 7, x + 7, y + 7), fill="#" + accent)
        dr.text((x - 33, 370), str(2020 + i), font=f, fill="#" + ink)
        dr.text((x - 22, y - 45), str(DATA[i]) + "%", font=bold, fill="#" + ink)
    im.save(out)


def build(plan, out, stress=False):
    d = plan["choices"]
    params = plan["parameters"]
    ident = plan["id"] + ("-stress" if stress else "")
    palette = plan["pagePalette"]["tokens"]
    accent = palette["action"].lstrip("#")
    background = palette["background"].lstrip("#")
    ink = palette["text"].lstrip("#")
    muted = palette["muted"].lstrip("#")
    on_accent = palette["onAction"].lstrip("#")
    heading, body = {
        "serif-sans": ("Georgia", "Calibri"),
        "sans-serif": ("Trebuchet MS", "Georgia"),
        "condensed-sans": ("Arial Narrow", "Calibri"),
        "mono-sans": ("Consolas", "Calibri"),
    }[d["typeRelationship"]]
    doc = Document()
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    margin = params["outerMarginIn"]
    s.left_margin = s.right_margin = Inches(margin)
    s.top_margin = s.bottom_margin = Inches(0.62)
    width = 8.5 - 2 * margin
    normal = doc.styles["Normal"]
    normal.font.name = body
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(ink)
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [
        ("Title", min(60, params["titleScalePt"] * 1.25)),
        ("Heading 1", 24),
        ("Heading 2", 15),
    ]:
        st = doc.styles[name]
        st.font.name = heading
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(ink)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0
        pp = st.element.find(qn("w:pPr"))
        if pp is not None:
            for edge in list(pp.findall(qn("w:pBdr"))):
                pp.remove(edge)
    # Word theme font attributes override explicit names unless removed.
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        fonts = doc.styles[style_name].element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for attr in list(fonts.attrib):
                if "theme" in attr.lower():
                    del fonts.attrib[attr]
    from lxml import etree

    page_art = etree.fromstring(
        f"""<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><v:rect id="DazzlerPageSurface" style="position:absolute;margin-left:0;margin-top:0;width:612pt;height:792pt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="#{background}" stroked="f"><v:fill color="#{background}"/><o:lock v:ext="edit" rotation="t"/></v:rect></w:pict>"""
    )
    bgp = s.header.add_paragraph()
    bgp.paragraph_format.space_after = Pt(0)
    bgp.paragraph_format.space_before = Pt(0)
    bgp.add_run()._r.append(page_art)
    hp = s.header.paragraphs[0]
    hp.text = "COMMON GROUND     /     URBAN SYSTEMS     /     STUDY 006"
    hp.style = doc.styles["Caption"]
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor.from_string(muted)
    if d["navigation"] == "running-band":
        shade(hp, accent)
        hp.runs[0].font.color.rgb = RGBColor.from_string(on_accent)
    elif d["navigation"] == "side-index":
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        border(hp, accent)
    fp = s.footer.paragraphs[0]
    fp.text = "SYNTHETIC DESIGN STUDY                                             "
    fp.runs[0].font.size = Pt(8)
    fp.runs[0].font.color.rgb = RGBColor.from_string(muted)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    art = out / (ident + "-art.png")
    plot = out / (ident + "-chart.png")
    illustration(art, d, accent, on_accent)
    chart(plot, accent, background, ink)
    paragraph(doc, "2020–2025  /  COVERAGE REVIEW", size=10, color=ink, bold=True)
    if d["opener"] == "framed":
        picture(
            doc, art, width, "Original abstract canopy motif; decorative, not a map"
        )
    title = paragraph(doc, TITLE, "Title")
    title.paragraph_format.space_after = Pt(18)
    for edge in list(title._p.get_or_add_pPr().findall(qn("w:pBdr"))):
        title._p.get_or_add_pPr().remove(edge)
    if d["alignment"] == "centered-title":
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if d["alignment"] == "offset":
        title.paragraph_format.left_indent = Inches(0.6)
    if d["opener"] == "side-rail":
        title.paragraph_format.left_indent = Inches(1.1)
        title.paragraph_format.right_indent = Inches(0.3)
    if d["colorRole"] in ["saturated-opener", "split-surface"]:
        shade(title, accent)
        for r in title.runs:
            r.font.color.rgb = RGBColor.from_string(on_accent)
    if d["opener"] == "evidence-led":
        picture(
            doc,
            plot,
            width,
            "Coverage by year: 2020 18%; 2021 21%; 2022 24%; 2023 29%; 2024 34%; 2025 42%. Zero baseline.",
        )
    paragraph(doc, INTRO, size=13 if d["density"] != "compact" else 11)
    if d["opener"] in ["split-field", "side-rail"]:
        picture(
            doc,
            art,
            width * 0.75 if d["opener"] == "side-rail" else width,
            "Original compositional motif; 42 filled units out of 100 when data-derived, otherwise illustrative",
        )
    p = paragraph(doc, "+24", size=96, color=accent, bold=True)
    p.paragraph_format.space_after = Pt(0)
    paragraph(doc, "Percentage Points of Additional Coverage", size=12, color=ink)
    paragraph(doc, FINDING)
    if d["opener"] == "type-led":
        picture(doc, art, width, "Original compositional motif; not a geographic map")
    paragraph(doc, LIMIT, size=9)
    paragraph(
        doc, "Evidence and Next Decisions", "Heading 1"
    ).paragraph_format.page_break_before = True
    paragraph(doc, "01   Read the Trend", "Heading 2")
    paragraph(doc, FINDING)
    if d["evidence"] in ["chart-led", "annotated"]:
        picture(
            doc,
            plot,
            width,
            "Six exact coverage values; zero baseline. Static chart, editable source table follows.",
        )
    if d["grid"] in ["two-column", "asymmetric", "modular"]:
        # Native newspaper columns for interior reading; no text boxes or screenshot pages.
        sec = doc.add_section(0)
        sec._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    paragraph(doc, "02   Compare the Source Values", "Heading 2")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Normal Table"
    for cell, text in zip(table.rows[0].cells, ["Year", "Coverage"]):
        cell.text = text
    for i, value in enumerate(DATA):
        cells = table.add_row().cells
        cells[0].text = str(2020 + i)
        cells[1].text = str(value) + "%"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(4)
                for r in p.runs:
                    r.font.size = Pt(10)
            if row_index == 0 or (d["table"] == "banded" and row_index % 2 == 0):
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), accent if row_index == 0 else background)
                cell._tc.get_or_add_tcPr().append(sh)
                if row_index == 0:
                    for r in cell.paragraphs[0].runs:
                        r.font.color.rgb = RGBColor.from_string(on_accent)
                        r.bold = True
            if d["table"] in ["rules", "minimal"]:
                border(cell.paragraphs[0], accent)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    paragraph(doc, "What the Series Can Tell Us", "Heading 2")
    paragraph(
        doc,
        "The annual increases are 3, 3, 5, 5 and 8 percentage points. The final interval has the largest increase. Each value describes a share of the same assumed study area; changing that boundary would make direct comparisons unreliable.",
    )
    if d["grid"] in ["two-column", "asymmetric", "modular"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
    paragraph(doc, "03   Decide What to Measure Next", "Heading 2")
    paragraph(doc, NEXT + (" " + NEXT if stress else ""))
    paragraph(doc, "Prepare the Next Review", "Heading 2")
    paragraph(
        doc,
        "Retain the geographic boundary and measurement definition used for each reporting year. Document any method change before comparing results. Where neighborhood figures are unavailable, mark the gap explicitly and keep the citywide result separate from local priorities.",
    )
    paragraph(doc, LIMIT, size=9)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = (
        "Original generative design exploration with synthetic data"
    )
    file = out / (ident + ".docx")
    doc.save(file)
    return {
        "id": ident,
        "choices": d,
        "parameters": params,
        "fonts": [heading, body],
        "accent": accent,
        "pageBackground": background,
        "colorChecks": plan["pagePalette"]["checks"],
        "file": file.name,
        "staticChart": True,
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--node", required=True)
    a = p.parse_args()
    a.out.mkdir(parents=True, exist_ok=False)
    req = {
        "purpose": "Create an expressive, readable two-page urban canopy review with editable evidence and clear next decisions",
        "content": ["prose", "data", "comparison"],
        "seed": "canopy-open-space-20261007",
        "count": 6,
    }
    request = a.out / "request.json"
    request.write_text(json.dumps(req), encoding="utf-8")
    raw = subprocess.check_output(
        [
            a.node,
            str(ROOT / "skills/dazzler-frontend/scripts/document-directions.mjs"),
            str(request),
        ],
        text=True,
    )
    plans = json.loads(raw)
    (a.out / "directions.json").write_text(raw, encoding="utf-8")
    results = [build(c, a.out) for c in plans["candidates"]]
    (a.out / "build.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
