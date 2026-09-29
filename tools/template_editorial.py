"""Shared semantic content rendered as distinct editorial HTML and editable Word."""

import html, re
from datetime import datetime, timezone
from pathlib import Path


def segments(text):
    pattern = r"\b(pilot|evidence|synthetic|decision|target|approved|review|owner|scope|retry|allergies|RSVP|fictional|at-least-once|acceptance|budget)\b"
    cursor = 0
    for match in list(re.finditer(pattern, str(text), re.I))[:3]:
        yield str(text)[cursor : match.start()], False
        yield match.group(), True
        cursor = match.end()
    yield str(text)[cursor:], False


def rich(text):
    return "".join(
        (
            '<strong class="keyword">' + html.escape(s) + "</strong>"
            if bold
            else html.escape(s)
        )
        for s, bold in segments(text)
    )


def render_html(d, path):
    e = lambda x: html.escape(str(x), quote=True)

    def block(b):
        k = b["type"]
        if k in ("p", "h"):
            return f'<{"h2" if k=="h" else "p"} class="{"section-title" if k=="h" else "prose"}">{rich(b["text"])}</{"h2" if k=="h" else "p"}>'
        if k == "metrics":
            return (
                '<div class="metrics">'
                + "".join(
                    '<div class="metric"><strong>'
                    + e(a)
                    + "</strong><span>"
                    + e(c)
                    + "</span></div>"
                    for a, c in b["values"]
                )
                + "</div>"
            )
        if k == "chart":
            return (
                '<figure class="chart"><figcaption>'
                + e(b["label"])
                + '</figcaption><img src="../charts/'
                + e(b["asset"])
                + '.svg" alt="'
                + e(b["label"] + "; exact values follow")
                + '"><details><summary>Read the chart data</summary><table><tbody>'
                + "".join(
                    "<tr><th>"
                    + e(a)
                    + "</th><td>"
                    + e(v)
                    + " "
                    + e(b["unit"])
                    + "</td></tr>"
                    for a, v in b["rows"]
                )
                + "</tbody></table></details></figure>"
            )
        if k == "callout":
            return (
                '<aside class="callout"><strong>'
                + e(b["label"])
                + "</strong><p>"
                + rich(b["text"])
                + "</p></aside>"
            )
        if k == "list":
            return (
                "<ul>"
                + "".join("<li>" + rich(t) + "</li>" for t in b["items"])
                + "</ul>"
            )
        if k == "code":
            return "<pre><code>" + e(b["text"]) + "</code></pre>"
        if d["id"] == "restaurant":
            return (
                '<div class="menu-course">'
                + "".join(
                    '<div class="dish"><div><h3>'
                    + e(row[0])
                    + "</h3><p>"
                    + e(row[1])
                    + "</p></div><strong>"
                    + e(row[-1])
                    + "</strong></div>"
                    for row in b["rows"]
                )
                + "</div>"
            )
        return (
            '<div class="table-wrap"><table><thead><tr>'
            + "".join('<th scope="col">' + e(x) + "</th>" for x in b["headers"])
            + "</tr></thead><tbody>"
            + "".join(
                "<tr>"
                + "".join(
                    "<"
                    + ('th scope="row"' if i == 0 else "td")
                    + ">"
                    + rich(v)
                    + "</"
                    + ("th" if i == 0 else "td")
                    + ">"
                    for i, v in enumerate(row)
                )
                + "</tr>"
                for row in b["rows"]
            )
            + "</tbody></table></div>"
        )

    pages = "".join(
        '<article class="sheet"><header class="folio"><span>'
        + e(d["tag"])
        + "</span><span>"
        + f'{i+1:02d} / {len(d["pages"]):02d}'
        + "</span></header><h1>"
        + e(pg["title"])
        + "</h1>"
        + (
            '<p class="subtitle"><em>' + e(d["subtitle"]) + "</em></p>"
            if i == 0
            else ""
        )
        + "".join(block(b) for b in pg["blocks"])
        + '<footer class="page-foot"><span>FICTIONAL EXAMPLE / '
        + e(d["id"].upper())
        + "</span><span>Dazzler · "
        + str(i + 1)
        + "</span></footer></article>"
        for i, pg in enumerate(d["pages"])
    )
    css = (Path(__file__).with_name("template_editorial.css")).read_text(
        encoding="utf-8"
    )
    if d.get("landscape"):
        css += "@page{size:letter landscape}"
    vars = f":root{{--accent:{d['accent']};--secondary:{d['secondary']};--bright:{d['bright']};--paper:{d['paper']};--ink:{d['ink']};--display:'{d['headingFont']}';}}"
    notes = (
        '<footer class="notes"><details><summary>Design notes and sample data</summary><p>'
        + e(d["capabilities"])
        + ". "
        + e(d["voice"])
        + '.</p><p>All organizations, people, dates and measurements are fictional. Replace sample material with verified project content. Original Dazzler design: Apache-2.0; local fonts retain their notices. Desktop Word fonts are referenced, not embedded.</p><a href="../data/'
        + d["id"]
        + '.json">Inspect the synthetic dataset</a></details></footer>'
    )
    path.write_text(
        '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>'
        + e(d["title"])
        + ' — Dazzler</title><link rel="stylesheet" href="../fonts/fonts.css"><style>'
        + vars
        + css
        + '</style></head><body class="doc-'
        + d["id"]
        + '"><nav class="tools"><a href="../index.html">Dazzler template gallery</a><button onclick="print()">Print / save PDF</button></nav><main>'
        + pages
        + "</main>"
        + notes
        + "</body></html>",
        encoding="utf-8",
    )


def render_docx(d, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    land = d.get("landscape", False)
    sec.page_width = Inches(11 if land else 8.5)
    sec.page_height = Inches(8.5 if land else 11)
    sec.top_margin = sec.bottom_margin = Inches(0.52)
    sec.left_margin = sec.right_margin = Inches(0.62)
    if land:
        sec.orientation = WD_ORIENT.LANDSCAPE
    accent = d["accent"] if d["id"] != "presentation" else "#393783"
    ink = "#172033"

    def color(run, value):
        run.font.color.rgb = RGBColor.from_string(value.lstrip("#"))

    def shade(node, value):
        el = OxmlElement("w:shd")
        el.set(qn("w:fill"), value.lstrip("#"))
        node.append(el)

    def runs(para, text, size=None):
        for value, bold in segments(text):
            run = para.add_run(value)
            run.bold = bold
            if bold:
                color(run, accent)
            if size:
                run.font.size = Pt(size)
        return para

    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn("w:rFonts")):
            for attr in list(fonts.attrib):
                if attr.endswith("Theme"):
                    del fonts.attrib[attr]
    title_size = {
        "fun": 43,
        "restaurant": 39,
        "presentation": 35,
        "family": 28,
        "marketing": 38,
        "technical": 30,
        "legal": 32,
    }.get(d["id"], 35)
    body_size = (
        9
        if d["id"] == "family"
        else (
            10
            if d["id"] == "presentation"
            else 10.5 if d["id"] in ("technical", "school", "business", "legal") else 11
        )
    )
    for name, size in [
        ("Normal", body_size),
        ("Title", title_size),
        ("Subtitle", 12),
        ("Heading 1", 15),
        ("Heading 2", 12),
        ("List Bullet", body_size),
    ]:
        st = doc.styles[name]
        st.font.name = (
            d["font"]
            if name in ("Title", "Heading 1") or d["id"] == "legal"
            else "Arial"
        )
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(
            (accent if name in ("Title", "Heading 1") else ink)[1:]
        )
        st.paragraph_format.space_after = Pt(
            4 if d["id"] in ("family", "presentation") else 7
        )
        st.paragraph_format.line_spacing = (
            1.0 if d["id"] in ("family", "presentation") else 1.1
        )
        if name in ("Title", "Heading 1"):
            st.paragraph_format.keep_with_next = True
            st.font.bold = d["id"] not in ("legal", "restaurant", "family", "school")
        if name == "Heading 1":
            st.paragraph_format.space_before = Pt(9)
    doc.core_properties.title = d["title"]
    doc.core_properties.author = "Dazzler / Jon Gosier"
    doc.core_properties.subject = d["use"]
    doc.core_properties.created = doc.core_properties.modified = datetime(
        2026, 9, 28, tzinfo=timezone.utc
    )
    footer = sec.footer.paragraphs[0]
    runs(
        footer, "FICTIONAL EXAMPLE  /  " + d["id"].upper() + "   •   Dazzler   /   ", 8
    )
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    def make_table(headers, rows, widths=None, metric=False):
        table = doc.add_table(rows=0, cols=len(headers))
        table.autofit = False
        width = sec.page_width - sec.left_margin - sec.right_margin
        fractions = widths or [100 / len(headers)] * len(headers)
        for col, fraction in zip(table.columns, fractions):
            col.width = int(width * fraction / 100)
        for ri, row in enumerate([headers, *rows]):
            cells = table.add_row().cells
            table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri == 0:
                table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ci, (cell, text) in enumerate(zip(cells, row)):
                cell.width = int(width * fractions[ci] / 100)
                props = cell._tc.get_or_add_tcPr()
                shade(
                    props, accent if ri == 0 else ("#F0F2F6" if ri % 2 else "#FFFFFF")
                )
                margin = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + side)
                    el.set(qn("w:w"), "45" if d["id"] == "family" else "85")
                    el.set(qn("w:type"), "dxa")
                    margin.append(el)
                props.append(margin)
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                if metric and ri == 0 and ci == 1:
                    shade(
                        props,
                        d["secondary"] if d["id"] != "presentation" else "#514287",
                    )
                run = para.add_run(str(text))
                run.font.name = "Arial"
                run.font.size = Pt(
                    (22 if d["id"] == "family" else 29)
                    if metric and ri == 0
                    else (
                        8.5
                        if d["id"] == "family"
                        else (
                            10
                            if d["id"] not in ("technical", "school", "business")
                            else 9.3
                        )
                    )
                )
                run.bold = ri == 0
                color(run, "#FFFFFF" if ri == 0 else ink)
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(0)
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.line_spacing = Pt(3)
        after.add_run().font.size = Pt(3)

    for i, pg in enumerate(d["pages"]):
        if i:
            doc.add_page_break()
        kicker = doc.add_paragraph()
        r = kicker.add_run(d["tag"] + "   /   " + str(i + 1).zfill(2))
        r.bold = True
        r.font.size = Pt(8)
        color(r, accent)
        title = doc.add_paragraph(pg["title"], "Title")
        if d["id"] == "presentation":
            shade(title._p.get_or_add_pPr(), "232447")
            for run in title.runs:
                color(run, "#E1F58B")
        if d["id"] == "restaurant":
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i == 0:
            para = doc.add_paragraph(d["subtitle"], "Subtitle")
            para.runs[0].italic = True
        for b in pg["blocks"]:
            k = b["type"]
            if k in ("p", "h"):
                runs(
                    doc.add_paragraph(style="Heading 1" if k == "h" else "Normal"),
                    b["text"],
                )
            elif k == "metrics":
                make_table(
                    [x[0] for x in b["values"]],
                    [[x[1] for x in b["values"]]],
                    metric=True,
                )
            elif k == "table":
                if d["id"] == "restaurant" and i == 0:
                    for name, description, price in b["rows"]:
                        para = doc.add_paragraph()
                        para.paragraph_format.space_before = Pt(12)
                        r = para.add_run(name + "   /   " + price)
                        r.font.name = "Georgia"
                        r.font.size = Pt(17)
                        color(r, accent)
                        para = doc.add_paragraph(description)
                        para.runs[0].italic = True
                        para.runs[0].font.size = Pt(11)
                else:
                    make_table(b["headers"], b["rows"], b.get("widths"))
            elif k == "chart":
                para = doc.add_paragraph(b["label"], "Heading 2")
                para.paragraph_format.keep_with_next = True
                doc.add_picture(
                    str(path.parents[1] / "charts" / f"{b['asset']}.png"),
                    width=Inches(6.2 if land else 6.65),
                )
            elif k == "list":
                for item in b["items"]:
                    runs(doc.add_paragraph(style="List Bullet"), item)
            elif k == "code":
                para = doc.add_paragraph()
                shade(para._p.get_or_add_pPr(), "EDF2F5")
                run = para.add_run(b["text"])
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            elif k == "callout":
                para = doc.add_paragraph()
                shade(
                    para._p.get_or_add_pPr(),
                    "F1EDF6" if d["id"] == "legal" else "EDF2F9",
                )
                para.paragraph_format.space_before = Pt(6)
                para.paragraph_format.space_after = Pt(8)
                r = para.add_run(b["label"] + "\n")
                r.bold = True
                r.font.size = Pt(9)
                color(r, accent)
                runs(para, b["text"])
        if i == len(d["pages"]) - 1:
            note = doc.add_paragraph()
            note.paragraph_format.space_before = Pt(7)
            r = note.add_run(
                "Design notes: "
                + d["capabilities"]
                + ". Original Dazzler template, Apache-2.0. All sample facts are fictional."
            )
            r.italic = True
            r.font.size = Pt(8)
            color(r, "#505968")
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal" and len(paragraph.text) > 80:
            section = doc.sections[0]
            available = (
                section.page_width - section.left_margin - section.right_margin
            ) / 914400
            paragraph.paragraph_format.right_indent = Inches(
                max(0, available - 60 * body_size * 0.5 / 72)
            )
    doc.save(path)
