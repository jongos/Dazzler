"""Optional native export integration check; run in a host with python-docx and python-pptx."""
import importlib.util
import json
from pathlib import Path
import sys
from docx import Document
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('project',ROOT/'skills/dazzler-frontend/scripts/project.py');project=importlib.util.module_from_spec(spec);spec.loader.exec_module(project)
out=Path(sys.argv[1]);system=project.read(out/'system/design-system.json')
content={'title':'Dazzler studio — illustrative brief','subtitle':'Export fixture, not client data','sections':[{'heading':'Purpose','body':'Build a welcoming, readable interface. Preserve the brand and prioritize the main task.\n\nThis example tests branded document and slide exports.'},{'heading':'Milestones','body':'The following table is illustrative.','table':[['Stage','Status'],['Design','Draft'],['Review','Scheduled']]},{'heading':'Long content','body':'A readable presentation carries one useful thought at a time. '*22}]}
font_root=Path(sys.argv[2]);content['fontExports']=[str(font_root/'young-serif'),str(font_root/'work-sans')]
for mode,extension in [('document','html'),('slides','html'),('docx','docx'),('pptx','pptx')]:
    result=project.export_document(content,system,out/(mode+'.'+extension),mode);print(json.dumps(result))
doc=Document(out/'docx.docx');assert len(doc.tables)==1;assert 'Purpose' in [p.text for p in doc.paragraphs];assert doc.styles['Normal'].font.name=='Work Sans'
prs=Presentation(out/'pptx.pptx');assert len(prs.slides)>len(content['sections'])+1;assert any(shape.has_table for slide in prs.slides for shape in slide.shapes)
assert 'data:font/' in (out/'document.html').read_text(encoding='utf-8');assert 'SIL OPEN FONT LICENSE' in (out/'document.html').read_text(encoding='utf-8')
print('HTML embedded fonts/notices and native DOCX/PPTX content, tables and continuation structure passed. Native visual layout remains unverified.')
