"""Apply recorded color decisions to editable documents and working interfaces."""

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RECORD = ROOT / "maintenance/gallery-production/v0.28.0/color-decisions.json"


def decision(ident):
    if not RECORD.exists():
        raise ValueError("Run gallery_color_exploration.mjs before authoring")
    return json.loads(RECORD.read_text(encoding="utf8"))[ident]


def contrast(a, b):
    def luminance(c):
        rgb = [int(c[i : i + 2], 16) / 255 for i in (1, 3, 5)]
        rgb = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in rgb]
        return sum(v * k for v, k in zip(rgb, (0.2126, 0.7152, 0.0722)))

    x, y = sorted((luminance(a), luminance(b)))
    return (y + 0.05) / (x + 0.05)


def ink(background):
    return max(("#000000", "#FFFFFF"), key=lambda c: contrast(c, background))


def apply(d):
    from fonts import load, shortlist

    catalog = load()["fonts"]
    family = next(f for f in catalog if f["name"] == d["font"])
    role = "heading" if "heading" in family["roles"] else family["roles"][0]
    matches = shortlist(
        catalog,
        role,
        "editorial",
        d["title"] + " " + d.get("lead", ""),
        400,
        False,
        False,
        [],
        None,
        family["id"],
    )
    if not matches:
        raise ValueError(
            "Selected display font does not cover the actual heading: " + d["id"]
        )
    d["fontSelection"] = {
        k: matches[0][k] for k in ("id", "css_family", "coverage_checked", "license")
    }
    r = decision(d["id"])
    t = r["tokens"]
    d["originalAccent"], d["originalPaper"] = d["accent"], d["paper"]
    accent = t["brand"] if contrast(t["brand"], t["background"]) >= 3 else t["action"]
    if d["id"] == "technical":
        accent = "#19DDE7"
    d.update(paper=t["background"], accent=accent, colorDirection=r["brief"])
    d["prompt"] = re.sub(
        r"Use (?:a cobalt decision spread|broadcast yellow and charcoal|a quiet blue field|ink blue, coral and cream).*?(?=\. )",
        "",
        d["prompt"],
    )
    d["prompt"] += " Color direction: " + r["brief"]
    return d


def css(d):
    t = decision(d["id"])["tokens"]
    bg, fg, a = t["background"], t["text"], t["action"]
    on = ink(a)
    rules = f"""
/* Selected engine roles; accent fills and labels have separate foregrounds. */
:root{{--paper:{bg};--accent:{a};--ink:{fg};--on-accent:{on};--signal:{t['accent']};--counter:{t['secondary']}}}
body{{background:var(--paper);color:var(--ink)}}
p,li,blockquote{{max-width:60ch}}
.top,footer,td,th,.event,.notes,.task,.product,.dish-row,.menu-controls,.ledger{{border-color:var(--accent)}}
button,input,select,textarea{{color:var(--ink);border-color:var(--accent)}}
input,select,textarea{{background:var(--paper)}}
button{{background:transparent}}
button[aria-pressed=true],.band,.schedule,thead,.beats section:nth-child(2){{background:var(--accent);color:var(--on-accent)}}
:focus-visible{{outline:3px solid var(--accent);outline-offset:5px}}
::selection{{background:var(--accent);color:var(--on-accent)}}
.kicker,.closing{{color:var(--ink)}}
.stat,.event,.reading-value{{color:var(--accent)}}
@media print{{body{{background:var(--paper);color:var(--ink)}}*{{print-color-adjust:exact;-webkit-print-color-adjust:exact}}}}
"""
    extras = {
        "professional": ".decision aside{background:#1547EF;color:white;padding:28px}.decision aside *{color:inherit}.decision h1{color:#1547EF}",
        "legal": ".legal-rail .event{border-left:7px solid #C52A12}.legal-rail h1{color:#101010}.legal-rail blockquote{border-color:#C52A12}",
        "business": ".proposal aside{background:#112E21;color:#C9FF32;padding:30px}.proposal aside *{color:inherit}.proposal{border-bottom:10px solid #112E21}",
        "fun": ".poster-date{color:#29104F}.schedule{background:#29104F;color:#FFFF6C}.poster h1{color:#161016}",
        "family": ".days .kicker{color:inherit}.days section:nth-child(1){background:#1648E8;color:white}.days section:nth-child(2){background:#FFDC29;color:#171717}.days section:nth-child(3){background:#FF784E;color:#171717}.packing{background:#D3F4FF;color:#082841}",
        "presentation": ".beat-no{color:#F4FF53}.beats section:nth-child(2){background:#F4FF53;color:#251367}.beats section:nth-child(2) .beat-no{color:#251367}",
        "school": ".notebook aside{background:#12512F;color:#FFFFFF;padding:28px;border:0}.notebook aside *{color:inherit}.notebook textarea{background:#FFFFFF;color:#171717}.notebook h1{color:#12512F}",
        "marketing": ".campaign{border-bottom:14px solid #171717}.campaign-mark{opacity:1;color:#111}.campaign h1{position:relative;z-index:1;max-width:9ch}.campaign-mark{right:-15px;top:18px;font-size:14rem}thead{background:#172BC9;color:white}@media(max-width:700px){.campaign-mark{display:none}}",
        "restaurant": ".menu,.menu-head{color:var(--ink)}.price,.menu-price{color:#E2FF89}h1{color:#E2FF89}",
        "technical": ".runbook{background:#19DDE7;color:#05212C}.runbook aside{color:#05212C}.runbook .kicker,.runbook .stat{color:#05212C}.band{background:#A6002E;color:white}.figure{color:#BBF7FF}.closing{background:#FFD5DD;color:#6B0015;padding:18px}",
        "webapp-workspace": ".notes{border-left:8px solid #BF3E08;padding-left:20px}.manuscript textarea{color:#151515}.chapters{border-top:8px solid #BF3E08;padding-top:18px}",
        "webapp-board": ".board-title h1{color:#111}.task button{background:#4521B8;color:white;border-color:#4521B8}.stage{color:#4521B8}",
        "webapp-settings": ".clock{background:#C3FCFF;color:#082457;border-color:#C3FCFF}.settings form{border-top:8px solid #C3FCFF}fieldset{border-color:#C3FCFF}select{background:#FFFFFF;color:#082457}",
        "data-revenue": ".chart-field{background:#161321;color:white}.bar i{background:#C596FF}.chart-field figcaption{color:white}button[aria-pressed=true]{background:#6822CF;color:white}",
        "data-operations": ".flow{background:#061F27;color:#65FFD6;border:2px solid #65FFD6}.flow button{color:#65FFD6;border-color:#65FFD6}.flow button[aria-pressed=true]{background:#65FFD6;color:#061F27}.reading-value{color:#65FFD6}.alert{background:#FFB75E;color:#241B08;border-color:#FFB75E}",
        "restaurant-fine-dining": ".course-number{color:#FFBD75}.vesper h1{color:#FFF5E3}.top,footer,.columns{border-color:#AB80C5}",
        "restaurant-cafe": ".ticket{background:#FFFFFF;color:#172A72;border-color:#1640BA}.counter .product button,button[aria-pressed=true]{background:#1634A6;color:white;border-color:#1634A6}.cafe h1{color:#111}.ticket button{color:#172A72}",
        "restaurant-reservations": ".room{background:#FFFFFF;border-color:#086D3C}.room button{background:#FFFFFF;color:#086D3C;border-color:#086D3C}.room button[aria-pressed=true]{background:#086D3C;color:white}.reservation form{background:#FFE0D7;color:#371C12;padding:28px}.reservation form input,.reservation form select{background:white;color:#171717}.reservation form button{background:#086D3C;color:white}.window-label{border-color:#086D3C}",
        "restaurant-menu": ".dish-number{color:#651278}.dish-row h2,.menu-head h1{color:#2C1037}.menu-controls{border-color:#2C1037}.menu-controls input{background:#FFFFFF;color:#2C1037}",
        "business-portal": ".portal{background:#1648E8;color:white;border:0;padding:30px}.portal .kicker{color:white}.ledger button{background:#085C3A;color:white;border-color:#085C3A}",
    }
    return rules + extras[d["id"]]


def restyle_word(doc, d):
    """Resolve ink against each native cell/page; never rasterize document content."""
    from docx.oxml.ns import qn
    from docx.oxml import parse_xml

    r = decision(d["id"])
    t = r["tokens"]
    background = t["background"]
    s = doc.sections[0]
    for pict in list(s.header._element.iter(qn("w:pict"))):
        pict.getparent().remove(pict)
    s.header.paragraphs[0].add_run()._r.append(
        parse_xml(
            f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:rect id="DazzlerPageSurface" style="position:absolute;margin-left:0;margin-top:0;width:{s.page_width/12700}pt;height:{s.page_height/12700}pt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="{background}" stroked="f"><v:fill color="{background}"/></v:rect></w:pict>'
        )
    )
    if d["id"] == "family":
        for cell, fill in zip(
            doc.tables[0].rows[0].cells, ["1648E8", "FFDC29", "FF784E"]
        ):
            cell._tc.get_or_add_tcPr().find(qn("w:shd")).set(qn("w:fill"), fill)
    if d["id"] == "presentation":
        doc.tables[0].rows[0].cells[1]._tc.get_or_add_tcPr().find(qn("w:shd")).set(
            qn("w:fill"), "F4FF53"
        )
    if d["id"] == "business":
        from docx.shared import Inches
        from docx.oxml import OxmlElement

        ledger = doc.tables[-1]._tbl
        headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
        split = doc.add_table(rows=1, cols=2)
        split.columns[0].width = Inches(4.6)
        split.columns[1].width = Inches(4.8)
        widths = [1512, 3240, 1656]
        for node, width in zip(ledger.find(qn("w:tblGrid")), widths):
            node.set(qn("w:w"), str(width))
        for row in ledger.findall(qn("w:tr")):
            for cell, width in zip(row.findall(qn("w:tc")), widths):
                cell.find("w:tcPr/w:tcW", namespaces=cell.nsmap).set(
                    qn("w:w"), str(width)
                )
        headings[0]._p.addprevious(split._tbl)
        for p in headings:
            following = p._p.getnext()
            split.cell(0, 0)._tc.append(p._p)
            if following is not None and following.tag == qn("w:p"):
                split.cell(0, 0)._tc.append(following)
        split.cell(0, 1)._tc.append(ledger)
        split.cell(0, 1)._tc.append(OxmlElement("w:p"))
    for element in doc.element.body.iter(qn("w:shd")):
        fill = element.get(qn("w:fill"), "")
        if fill.upper() == "FFFFFF" and r["mode"] == "dark":
            element.set(qn("w:fill"), background[1:])
    for run in doc.element.body.iter(qn("w:r")):
        surface = background
        parent = run.getparent()
        while parent is not None and parent.tag != qn("w:body"):
            if parent.tag == qn("w:tc"):
                shd = parent.find("w:tcPr/w:shd", namespaces=parent.nsmap)
                if shd is not None and re.fullmatch(
                    r"[0-9A-Fa-f]{6}", shd.get(qn("w:fill"), "")
                ):
                    surface = "#" + shd.get(qn("w:fill"))
                break
            parent = parent.getparent()
        prop = run.get_or_add_rPr()
        color = prop.find(qn("w:color"))
        if color is None:
            from docx.oxml import OxmlElement

            color = OxmlElement("w:color")
            prop.append(color)
        old = "#" + color.get(qn("w:val"), "17252B")
        preferred = d["accent"] if old.upper() == d["accent"].upper() else t["text"]
        color.set(
            qn("w:val"),
            (preferred if contrast(preferred, surface) >= 4.5 else ink(surface))[1:],
        )
    for color in s.footer._element.iter(qn("w:color")):
        color.set(qn("w:val"), t["text"][1:])
    if d["id"] == "technical":
        for p in doc.paragraphs:
            if p.text.startswith("STOP:"):
                from docx.shared import RGBColor

                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string("FF8DA0")
