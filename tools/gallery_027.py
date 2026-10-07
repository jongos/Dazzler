"""Reproduce the independently authored October gallery, not the legacy preset builder.
Briefs, compositions and copy were authored with the current Dazzler workflow.
Staging is intentional: promote only after browser/native render review.
"""

from pathlib import Path
import json, html, shutil, subprocess, hashlib, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills/dazzler-frontend"
STAGE = ROOT / "maintenance/gallery-production/v0.28.0"
OUT = STAGE / "artifacts"
sys.path.insert(0, str(SKILL / "scripts"))
from headings import heading

DOCS = [
    dict(
        id="professional",
        title="Harbor Exchange Launch Decision",
        short="The Launch Decision",
        tag="NORTHLINE / OPERATING REVIEW",
        lead="Release the exchange in two stages. Keep the settlement hold until reconciliation is complete.",
        accent="#123A8C",
        paper="#EEF3FF",
        font="Archivo",
        word="Arial",
        layout="decision",
        stat="92%",
        label="of pilot transactions reconciled",
        intro="The six-week pilot processed 1,250 synthetic transactions. Reconciliation matched 1,150; the remaining 100 need manual review before settlement can be enabled.",
        sections=[
            (
                "Approve the Limited Release",
                "Open account setup and order tracking on 14 July. Keep settlement disabled until the exception queue is empty and Finance signs the control record.",
            ),
            (
                "What Changed",
                "A single owner now resolves each exception. The support team has a tested escalation route, and the release can be rolled back without losing the review history.",
            ),
            (
                "Decision Owners",
                "Mara Chen owns reconciliation; Omar Diaz owns rollback readiness. Both report to the release lead at 09:00 on launch day.",
            ),
        ],
        headers=["Gate", "Evidence", "State"],
        rows=[
            ["Account setup", "240 completed test journeys", "Ready"],
            ["Tracking", "98% delivered within 2 seconds", "Ready"],
            ["Settlement", "100 unmatched records", "Hold"],
        ],
        closing="Decision due 10 July · Approve the limited scope, not settlement.",
        prompt="Create a fictional executive launch-decision brief for an operations committee. Make the decision, unresolved gate and accountable owners unmistakable. Use a cobalt decision spread with precise numbers and a compact gate register; avoid a generic KPI dashboard.",
    ),
    dict(
        id="legal",
        title="Orchard Lane Notice Review",
        short="Notice and Evidence",
        tag="MATTER 026 / INTERNAL REVIEW",
        lead="The record supports delivery. It does not yet establish that the recipient accepted the revised terms.",
        accent="#69283E",
        paper="#FAF5EC",
        font="Libre Baskerville",
        word="Georgia",
        layout="chronology",
        stat="03",
        label="events establish the sequence",
        intro="This fictional matter note separates the available correspondence from the questions counsel must resolve. It contains no jurisdiction-specific legal conclusion.",
        sections=[
            (
                "Issue Presented",
                "Which version of the terms accompanied the notice, and what evidence records the recipient's response? Review the document trail before drawing a conclusion.",
            ),
            (
                "Record Gaps",
                "The dispatch log identifies an attachment but does not preserve its hash. A reply refers to revised dates without identifying a signed amendment.",
            ),
            (
                "Next Review",
                "Obtain the original attachment and complete message headers. Keep the current chronology provisional until the underlying files are reconciled.",
            ),
        ],
        headers=["Date", "Record", "Significance"],
        rows=[
            ["04 May", "Notice dispatched", "Delivery initiated"],
            ["06 May", "Receipt acknowledged", "Receipt only"],
            ["09 May", "Dates discussed", "Acceptance unresolved"],
        ],
        closing="Prepared for counsel · Fictional training material, not legal advice.",
        prompt="Design a fictional legal evidence-review memorandum. Use a narrow plum chronology and a quiet cream reading field with serif hierarchy. Distinguish documented events from unresolved conclusions; no legal advice, invented authorities or decorative legal icons.",
    ),
    dict(
        id="business",
        title="Juniper Studio Service Proposal",
        short="A Better Booking Journey",
        tag="JUNIPER / SCOPE 014",
        lead="Replace the fragmented booking journey with one accessible flow, delivered in three reviewable stages.",
        accent="#145A45",
        paper="#F1F6E9",
        font="Young Serif",
        word="Georgia",
        layout="proposal",
        stat="$24,000",
        label="fixed fee for the defined scope",
        intro="The fictional engagement covers research synthesis, interaction design and a tested frontend prototype for a neighborhood arts venue. Production integration is a separate engagement.",
        sections=[
            (
                "Included",
                "Two stakeholder workshops, six core screens, responsive specifications, accessible component states and a documented prototype handoff.",
            ),
            (
                "Working Together",
                "One named decision owner consolidates feedback within three working days. Each stage ends with a review, not an implied approval.",
            ),
            (
                "Commercial Terms",
                "Invoices follow the milestones below. The estimate excludes taxes, third-party subscriptions, production hosting and changes outside the agreed six screens.",
            ),
        ],
        headers=["Stage", "Deliverable", "Fee"],
        rows=[
            ["01 Discover", "Journey and success measures", "$6,000"],
            ["02 Design", "Six screens and component states", "$10,000"],
            ["03 Validate", "Prototype and handoff", "$8,000"],
        ],
        closing="Offer valid for 30 days · Start date follows written acceptance.",
        prompt="Make a fictional creative-service proposal for an arts venue. Give the scope and fixed fee strong editorial hierarchy. Compose an oversized fee beside a three-stage delivery ledger, with forest green, pale lime and genuine serif italics; keep exclusions legible.",
    ),
    dict(
        id="fun",
        title="Night Market on Alder Street",
        short="Meet After Sunset",
        tag="ALDER STREET / SATURDAY 17 JULY",
        lead="An evening of food, records and small discoveries. Bring a friend and leave room for dessert.",
        accent="#64204C",
        paper="#FFE9A8",
        font="Oswald",
        word="Arial",
        layout="poster",
        stat="17 JUL",
        label="17:00–22:00 · free entry",
        intro="This fictional neighborhood night market fills the pedestrian block between Pine and Alder. No tickets are required; individual stalls set their own prices.",
        sections=[
            (
                "Make an Evening of It",
                "Start at the community table, browse the makers and catch the sunset set. Seating is shared and the route is step-free.",
            ),
            (
                "Come Prepared",
                "Bring a reusable bottle. Water refill points are beside the information tent; quiet seating is at the library end of the block.",
            ),
            (
                "Weather Plan",
                "If heavy rain is forecast, the event moves to Alder Hall. Check the information board before traveling.",
            ),
        ],
        headers=["Time", "Find", "Where"],
        rows=[
            ["17:00", "Food Stalls Open", "East end"],
            ["18:30", "Listening Session", "Record tent"],
            ["20:00", "Sunset Set", "Community stage"],
        ],
        closing="Fictional event · No real booking or attendance service.",
        prompt="Create a joyful fictional neighborhood night-market program. Use a typographic date poster, aubergine ink on warm yellow, a bold schedule ribbon and generous negative space. Make access, quiet seating and the rain plan as useful as the event headline.",
    ),
    dict(
        id="family",
        title="River House Weekend Plan",
        short="One Weekend Together",
        tag="THE RIVER HOUSE / 23–25 JULY",
        lead="A shared plan with room to change it. The only fixed appointment is lunch with Nana on Saturday.",
        accent="#224A74",
        paper="#FFF7E6",
        font="Work Sans",
        word="Arial",
        layout="planner",
        stat="3 DAYS",
        label="one shared list · fewer messages",
        intro="A fictional family plan for two adults and two children. Keep food, travel and quiet time visible; check items as they are packed rather than assuming they are done.",
        sections=[
            (
                "Pack Together",
                "Rain jackets, refillable bottles, medication, chargers and one book each. Keep medication with the adult responsible for it.",
            ),
            (
                "Food Plan",
                "Bring Friday supper. Buy breakfast ingredients locally. Confirm dietary needs before Saturday lunch; the sample plan does not assume medical requirements.",
            ),
            (
                "If Plans Change",
                "Use the shared notice space for the new meeting time. Agree on one pickup point before anyone leaves the house.",
            ),
        ],
        headers=["Day", "Anchor Plan", "Leave Open"],
        rows=[
            ["Friday", "Arrive before 18:00", "Games after supper"],
            ["Saturday", "Lunch at 12:30", "Walk or quiet time"],
            ["Sunday", "Pack by 11:00", "One last river stop"],
        ],
        closing="Fictional household · Adapt the plan to your family and access needs.",
        prompt="Design a warm practical family weekend planner. Use three day columns, a packing checklist and a shared-notice strip. It should feel like a considered fridge sheet in cream, navy and coral, not a corporate report or a childish worksheet.",
    ),
    dict(
        id="presentation",
        title="Library After Hours Pilot",
        short="Open Later for More People",
        tag="CIVIC LAB / THREE MINUTE BRIEF",
        lead="Test one late opening each week before committing to a permanent change.",
        accent="#482884",
        paper="#F5EFFF",
        font="Archivo",
        word="Arial",
        layout="storyboard",
        stat="6 WEEKS",
        label="a bounded pilot with a clear review point",
        intro="This fictional presentation handout proposes Thursday openings until 21:00. Attendance, staffing and quiet-study demand will determine the next decision.",
        sections=[
            (
                "The Need",
                "Evening visitors need access after work and classes. The pilot tests demand rather than assuming every weekday should open later.",
            ),
            (
                "The Experiment",
                "Open on six Thursdays. Keep a quiet study zone, a staffed help point and a clear closing routine. Record occupancy every hour.",
            ),
            (
                "The Decision",
                "Continue only if demand is sustained and the staffing plan is workable. Publish the measured results, including low-attendance evenings.",
            ),
        ],
        headers=["Measure", "Collection", "Review"],
        rows=[
            ["Visitors", "Hourly count", "Demand by hour"],
            ["Quiet seats", "Occupied / available", "Capacity pressure"],
            ["Staff hours", "Scheduled and actual", "Operating effort"],
        ],
        closing="Pilot proposal only · Synthetic example, not a funded program.",
        prompt="Create a landscape presentation handout for a fictional late-opening library pilot. Organize it as three visual beats: need, experiment, decision. Use large purple typographic numerals, concise evidence, and an evaluation strip rather than a report page or a fake slide deck.",
    ),
    dict(
        id="school",
        title="Rain Garden Observation Notebook",
        short="Where Does the Water Go",
        tag="FIELD SCIENCE / YEAR 8",
        lead="Compare drainage time in three soil samples. Record what happened before deciding why.",
        accent="#28592E",
        paper="#F3F1D9",
        font="Young Serif",
        word="Georgia",
        layout="notebook",
        stat="240 mL",
        label="the same water volume for every sample",
        intro="In this fictional classroom investigation, three equal containers hold sand, garden soil and clay-rich soil. Each receives the same water volume; students time the first 100 mL collected.",
        sections=[
            (
                "Keep the Test Fair",
                "Use the same container size, water volume and timing method. Repeat each sample three times. Describe any spills or delayed starts.",
            ),
            (
                "Read the Pattern",
                "Sand drained fastest in the sample results. The test measures this setup only; it does not establish how an entire garden will behave.",
            ),
            (
                "Ask a Better Question",
                "How would compaction change the result? Plan one change at a time and state which measurement would answer your question.",
            ),
        ],
        headers=["Sample", "Trial Times", "Median"],
        rows=[
            ["Sand", "38 / 41 / 40 sec", "40 sec"],
            ["Garden soil", "84 / 90 / 87 sec", "87 sec"],
            ["Clay-rich soil", "165 / 172 / 168 sec", "168 sec"],
        ],
        closing="Synthetic classroom observations · Follow your teacher's safety instructions.",
        prompt="Make a fictional Year 8 science field notebook on soil drainage. Use botanical green, warm paper, ruled observations and a labeled experimental diagram. Keep methods, measurements and limitations distinct; invite a next question rather than dressing a generic report in green.",
    ),
    dict(
        id="marketing",
        title="Second Life Repair Campaign",
        short="Keep the Good Things",
        tag="SECOND LIFE / CAMPAIGN EDITION 01",
        lead="Make repair feel like a first choice. Show the work, explain the process and make booking simple.",
        accent="#9B301C",
        paper="#FFF0E5",
        font="Oswald",
        word="Arial",
        layout="editorial",
        stat="120",
        label="illustrative repair slots over four weeks",
        intro="A fictional campaign for a local repair collective. The offer covers a 20-minute assessment; repair price and turnaround are confirmed after inspection.",
        sections=[
            (
                "The Story",
                "Lead with a well-used object and its repair, not an unsupported sustainability promise. Explain what can be fixed and what cannot.",
            ),
            (
                "The Offer",
                "Book an assessment, bring the item and receive a written estimate. No repair begins without approval. Parts may affect timing.",
            ),
            (
                "Measure the Journey",
                "Track assessment bookings, attended appointments and approved repairs separately. Do not present bookings as completed repairs.",
            ),
        ],
        headers=["Channel", "Message", "Action"],
        rows=[
            ["Window poster", "Keep the Good Things", "Scan for times"],
            ["Email", "See a repair in progress", "Book assessment"],
            ["Workshop", "Meet the repair team", "Bring one item"],
        ],
        closing="Fictional campaign · Capacity figures are planning assumptions.",
        prompt="Build a bold fictional repair-campaign editorial. Use a large condensed headline, rust and peach color fields, a diagram of the repair journey and a channel ledger. Avoid generic marketing cards and unsupported environmental claims.",
    ),
    dict(
        id="restaurant",
        title="Morrow Seasonal Supper Menu",
        short="An Evening at Morrow",
        tag="MORROW / LATE SUMMER",
        lead="A short menu built around what is ready now. Choose a few dishes and share the table.",
        accent="#713541",
        paper="#F9F1DF",
        font="Libre Baskerville",
        word="Georgia",
        layout="menu",
        stat="18:00",
        label="first seating · Wednesday to Sunday",
        intro="A fictional neighborhood supper room. Prices are sample USD amounts before tax. Ask the team about ingredients and cross-contact; dietary labels are not safety guarantees.",
        sections=[
            (
                "To Begin",
                "Charred peach, tomato and basil · 12\nWhite bean broth, lemon and herbs · 10",
            ),
            (
                "For the Table",
                "Roast squash, farro and sage · 24\nPan-roasted trout, potatoes and greens · 29",
            ),
            (
                "Something Sweet",
                "Plum tart with cultured cream · 11\nOlive oil cake and poached pear · 10",
            ),
        ],
        headers=["Pairing", "Glass", "Character"],
        rows=[
            ["Orchard Spritz", "9", "Apple / rosemary / soda"],
            ["House White", "12", "Dry / citrus / mineral"],
            ["Garden Tonic", "8", "Cucumber / mint / tonic"],
        ],
        closing="Sample menu · Tell the team about allergies before ordering.",
        prompt="Design an original seasonal supper menu for fictional Morrow. Use a wine-colored vertical wordmark, elegant serif course hierarchy, aligned prices and a botanical line motif. Give ingredient and allergy information readable space; avoid a report shell or stock-food hero.",
    ),
    dict(
        id="technical",
        title="Relay Queue Recovery Runbook",
        short="Restore the Queue Safely",
        tag="RELAY / OPERATIONS RUNBOOK 07",
        lead="Pause delivery before replaying events. Confirm the checkpoint, then resume in a controlled batch.",
        accent="#07556A",
        paper="#EAF5F5",
        font="Office Code Pro",
        word="Consolas",
        layout="runbook",
        stat="50",
        label="events in the first recovery batch",
        intro="This fictional runbook handles a stalled delivery queue. It demonstrates a safe sequence; adapt commands and thresholds to the actual system before operational use.",
        sections=[
            (
                "01 Stabilize",
                "Pause the worker and record the last acknowledged event ID. Confirm that the consumer is healthy before attempting a replay.",
            ),
            (
                "02 Recover",
                "Replay the next 50 events from the checkpoint. Check duplicate handling and compare delivered IDs with acknowledgements.",
            ),
            (
                "03 Verify",
                "Resume only when the batch reconciles and error rate remains below the agreed threshold. If it fails, pause again and retain the evidence.",
            ),
        ],
        headers=["Signal", "Proceed", "Stop"],
        rows=[
            ["Acknowledgements", "50 of 50 matched", "Any missing ID"],
            ["Duplicates", "Handled idempotently", "Side effect repeated"],
            ["Consumer health", "Checks passing", "Timeout or error"],
        ],
        closing="Synthetic operational example · Commands require environment-specific review.",
        prompt="Create a fictional technical recovery runbook. Make the pause-replay-verify sequence the composition, with monospaced step labels, a cyan evidence rail and a high-contrast stop condition. Preserve exact operational language and provide a rollback branch, not a generic dashboard.",
    ),
]


from gallery_color_design import apply as apply_color, css as color_css, restyle_word

DOCS = [apply_color(d) for d in DOCS]


def esc(x):
    return html.escape(str(x), quote=True)


def table(d):
    return (
        '<p class="table-hint">Scroll the table sideways to see all columns.</p><div class="table-scroll" tabindex="0" role="region" aria-label="'
        + esc(d["short"])
        + ' data"><table><thead><tr>'
        + "".join('<th scope="col">' + esc(x) + "</th>" for x in d["headers"])
        + "</tr></thead><tbody>"
        + "".join(
            "<tr>"
            + "".join(
                ('<th scope="row">' if i == 0 else "<td>")
                + esc(x)
                + ("</th>" if i == 0 else "</td>")
                for i, x in enumerate(row)
            )
            + "</tr>"
            for row in d["rows"]
        )
        + "</tbody></table></div>"
    )


def sections(d):
    return "".join(
        '<section class="reading"><h2>'
        + esc(heading(h))
        + "</h2><p>"
        + esc(p).replace("\n", "<br>")
        + "</p></section>"
        for h, p in d["sections"]
    )


def motif(d):
    # Original vector notation describes the content; no borrowed illustration.
    labels = {
        "notebook": ["Equal Volume", "Same Container", "Timed Drainage"],
        "runbook": ["Pause", "Replay", "Verify"],
        "editorial": ["Assess", "Estimate", "Repair"],
    }.get(d["layout"], ["Observe", "Decide", "Act"])
    return (
        '<svg viewBox="0 0 720 150" role="img" aria-label="'
        + esc(" then ".join(labels))
        + '"><path d="M70 65H650" stroke="currentColor" stroke-width="2"/>'
        + "".join(
            f'<circle cx="{110+i*250}" cy="65" r="35" fill="var(--paper)" stroke="currentColor" stroke-width="3"/><text x="{110+i*250}" y="74" text-anchor="middle" fill="currentColor" font-size="25">{i+1}</text><text x="{110+i*250}" y="135" text-anchor="middle" fill="currentColor" font-size="18">{esc(t)}</text>'
            for i, t in enumerate(labels)
        )
        + "</svg>"
    )


BASE = """*{box-sizing:border-box}html{scroll-behavior:auto}body{margin:0;background:var(--paper);color:#17252b;font:17px/1.6 "Work Sans",Arial,sans-serif}h1,h2,h3,p{margin-top:0}h1,h2,h3{line-height:1.1}h1{font:600 clamp(2.7rem,6vw,6.4rem)/1.02 var(--display);letter-spacing:-.04em;margin-bottom:24px}h2{font-size:1.65rem;margin:0 0 16px}h3{font-size:1.2rem}p{max-width:60ch}.kicker{font:600 12px/1.4 "Work Sans";letter-spacing:.14em;text-transform:uppercase}.lead{font-size:clamp(1.2rem,2vw,1.7rem);line-height:1.4}a{color:inherit;text-underline-offset:4px}button,input,select{font:inherit}button{cursor:pointer;padding:12px 20px;min-height:44px;border:1px solid currentColor;background:transparent;color:inherit}button:hover{background:#ffffff30}input,select{padding:12px;max-width:100%;border:1px solid #637079;background:#fff;color:#17252b}label{display:block}a:focus-visible,button:focus-visible,input:focus-visible,select:focus-visible,[tabindex]:focus-visible,summary:focus-visible{outline:3px solid #AE471D;outline-offset:4px}.top{display:flex;justify-content:space-between;align-items:center;gap:16px;max-width:1320px;margin:auto;padding:22px 5vw;border-bottom:1px solid #17252b40}.top a{font-size:14px}.canvas{max-width:1320px;margin:auto;padding:64px 5vw}.stat{font:600 clamp(3rem,7vw,7rem)/1 var(--display);letter-spacing:-.05em}.stat-label{font-size:14px;max-width:25ch}.reading{padding:24px 0}.reading p{white-space:normal}.table-hint{display:none}.table-scroll{overflow-x:auto;margin:28px 0}table{width:100%;border-collapse:collapse;font-size:14px;text-align:left}th,td{padding:14px 12px;border-bottom:1px solid #17252b40;vertical-align:top;min-width:100px}th{font-weight:600}thead{background:var(--accent);color:white}td{font-variant-numeric:tabular-nums}footer{padding:24px 5vw;border-top:1px solid #17252b40;font-size:12px;max-width:1320px;margin:auto}.closing{font-weight:600;font-size:15px;padding:22px 0}.canvas>*,.columns>*,.decision>*,.proposal>*,.beats>*,.notebook>*,.runbook>*{min-width:0}.columns{display:grid;grid-template-columns:1fr 1fr;gap:48px}.band{padding:28px;background:var(--accent);color:white}.band p{margin-bottom:0}.figure{max-width:800px;color:var(--accent);margin:32px 0}svg{display:block;width:100%;height:auto}.actions{display:flex;gap:12px;flex-wrap:wrap}.small{font-size:13px}.status{min-height:1.6em;font-weight:600}details{padding:14px 0}summary{cursor:pointer;min-height:32px}blockquote{font:400 1.5rem/1.4 var(--display);margin:20px 0;padding-left:20px;border-left:3px solid var(--accent)}@media(max-width:700px){.table-hint{display:block;font-size:13px;margin:18px 0 -12px}.canvas{padding:36px 22px}.top{padding:16px 22px;flex-wrap:wrap}.columns{grid-template-columns:1fr;gap:16px}h1{overflow-wrap:anywhere}table{min-width:520px}.stat{font-size:4rem}}@media print{.top,.actions,button{display:none}body{background:white}.canvas{padding:0;max-width:none}h1{font-size:36pt}p{font-size:11pt}section,table,svg{break-inside:avoid}footer{font-size:8pt}.table-scroll{overflow:visible}*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}"""


def document_html(d):
    title = (
        '<p class="kicker">'
        + esc(d["tag"])
        + "</p><h1>"
        + esc(heading(d["short"]))
        + '</h1><p class="lead">'
        + esc(d["lead"])
        + "</p>"
    )
    stat = (
        '<div class="stat">'
        + esc(d["stat"])
        + '</div><p class="stat-label">'
        + esc(d["label"])
        + "</p>"
    )
    intro = "<p>" + esc(d["intro"]) + "</p>"
    sec = sections(d)
    tab = table(d)
    fig = '<div class="figure">' + motif(d) + "</div>"
    k = d["layout"]
    css = ""
    if k == "decision":
        body = (
            '<div class="decision"><header>'
            + title
            + '</header><aside class="band">'
            + stat
            + "<p>SETTLEMENT REMAINS ON HOLD</p></aside><div>"
            + intro
            + sec
            + '</div><div class="gate"><h2>Release Gates</h2>'
            + tab
            + "</div></div>"
        )
        css = ".decision{display:grid;grid-template-columns:1.4fr 1fr;gap:42px}.decision aside{display:flex;flex-direction:column;justify-content:center}.decision h1{font-size:clamp(3rem,6vw,6rem)}.gate{padding-top:24px}@media(max-width:700px){.decision{grid-template-columns:1fr}}"
    elif k == "chronology":
        body = (
            '<div class="legal-rail"><aside><p class="kicker">RECORD SEQUENCE</p>'
            + "".join(
                '<div class="event"><strong>'
                + esc(r[0])
                + "</strong><p>"
                + esc(r[1])
                + "</p></div>"
                for r in d["rows"]
            )
            + "</aside><article>"
            + title
            + intro
            + sec
            + "<details><summary>Open the Evidence Register</summary>"
            + tab
            + "</details></article></div>"
        )
        css = ".legal-rail{display:grid;grid-template-columns:200px 1fr;gap:70px}.legal-rail article{max-width:720px}.event{border-left:2px solid var(--accent);padding:20px;color:var(--accent)}.legal-rail h1{font-size:clamp(3rem,5vw,5rem)}.reading h2{font-family:var(--display);font-weight:400}@media(max-width:700px){.legal-rail{display:flex;flex-direction:column-reverse;gap:24px}}"
    elif k == "proposal":
        body = (
            '<header class="proposal">'
            + title
            + "<aside>"
            + stat
            + '</aside></header><div class="columns"><div>'
            + intro
            + sec
            + "</div><div><h2>The Delivery Ledger</h2>"
            + tab
            + "<blockquote>One decision owner. Three review points. A defined finish.</blockquote></div></div>"
        )
        css = ".proposal{display:grid;grid-template-columns:1.3fr 1fr;gap:40px;border-bottom:3px solid var(--accent);padding-bottom:40px;margin-bottom:42px}.proposal aside{grid-column:2;grid-row:1/4;align-self:center;color:var(--accent)}@media(max-width:700px){.proposal{display:block}.proposal aside{margin-top:32px}}"
    elif k == "poster":
        body = (
            '<header class="poster"><p class="kicker">'
            + esc(d["tag"])
            + '</p><div class="poster-date">17<br>JUL</div><h1>'
            + esc(heading(d["short"]))
            + '</h1><p class="lead">'
            + esc(d["lead"])
            + '</p></header><div class="schedule">'
            + "".join(
                "<div><strong>"
                + esc(r[0])
                + "</strong><h2>"
                + esc(heading(r[1]))
                + "</h2><p>"
                + esc(r[2])
                + "</p></div>"
                for r in d["rows"]
            )
            + '</div><div class="columns">'
            + intro
            + "<div>"
            + sec
            + "</div></div>"
        )
        css = ".poster{display:grid;grid-template-columns:1fr 1fr;gap:24px}.poster-date{grid-column:2;grid-row:1/4;font:700 clamp(7rem,18vw,17rem)/1.05 Oswald;color:var(--accent);text-align:right}.poster h1{font-size:clamp(4rem,8vw,8rem);text-transform:uppercase}.schedule{display:grid;grid-template-columns:repeat(3,1fr);background:var(--accent);color:white;padding:32px;gap:24px;margin:50px 0}.schedule strong{font-size:2rem}@media(max-width:700px){.poster{display:block}.poster-date{text-align:left;float:right;font-size:7rem}.schedule{grid-template-columns:1fr}}"
    elif k == "planner":
        body = (
            "<header>"
            + title
            + '</header><div class="days">'
            + "".join(
                '<section><p class="kicker">'
                + esc(r[0])
                + "</p><h2>"
                + esc(heading(r[1]))
                + "</h2><p>"
                + esc(r[2])
                + '</p><label><input type="checkbox"> Plan Confirmed</label></section>'
                for r in d["rows"]
            )
            + '</div><div class="columns"><div>'
            + intro
            + sec
            + '</div><aside class="packing"><h2>Ready by the Door</h2>'
            + "".join(
                '<label><input type="checkbox"> ' + x + "</label>"
                for x in [
                    "Rain jackets",
                    "Refillable bottles",
                    "Medication with responsible adult",
                    "Chargers",
                    "One book each",
                ]
            )
            + '<p class="small">Checks stay in this page only.</p></aside></div>'
        )
        css = ".days{display:grid;grid-template-columns:repeat(3,1fr);gap:0;margin:36px 0;border:2px solid var(--accent)}.days section{padding:26px;border-right:1px solid var(--accent)}.days h2{font-size:1.5rem}.packing{padding:32px;background:#F8D6C4;align-self:start}.packing label{padding:12px 0;border-bottom:1px solid #17252b40}input[type=checkbox]{width:20px;height:20px;vertical-align:middle;margin-right:12px}@media(max-width:700px){.days{grid-template-columns:1fr}.days section{border-bottom:1px solid var(--accent)}}"
    elif k == "storyboard":
        body = (
            "<header>"
            + title
            + '</header><div class="beats">'
            + "".join(
                '<section><div class="beat-no">0'
                + str(i + 1)
                + "</div><h2>"
                + esc(heading(h))
                + "</h2><p>"
                + esc(p)
                + "</p></section>"
                for i, (h, p) in enumerate(d["sections"])
            )
            + '</div><div class="columns"><aside>'
            + stat
            + intro
            + "</aside><div><h2>What We Will Measure</h2>"
            + tab
            + "</div></div>"
        )
        css = ".beats{display:grid;grid-template-columns:repeat(3,1fr);gap:28px;margin:44px 0}.beats section{padding:24px 0;border-top:3px solid var(--accent)}.beat-no{font:700 7rem/1 Archivo;color:var(--accent);margin-bottom:24px}.beats section:nth-child(2){background:var(--accent);color:white;padding:24px}.beats section:nth-child(2) .beat-no{color:white}@media(max-width:700px){.beats{grid-template-columns:1fr}}"
    elif k == "notebook":
        body = (
            '<div class="notebook"><header>'
            + title
            + intro
            + "</header><aside>"
            + stat
            + fig
            + "</aside><article>"
            + sec
            + "</article><section><h2>Observation Record</h2>"
            + tab
            + '<label for="question">Your Next Question</label><textarea id="question" rows="5" placeholder="What would you change in the next test?"></textarea><p class="small">Local notes are not submitted or saved.</p></section></div>'
        )
        css = ".notebook{display:grid;grid-template-columns:1.1fr 1fr;gap:48px}.notebook aside{border-left:2px solid var(--accent);padding-left:32px;color:var(--accent)}textarea{font:inherit;width:100%;padding:16px;border:1px solid var(--accent);background:repeating-linear-gradient(#fff 0 31px,#c6cbbb 31px 32px);line-height:32px}@media(max-width:700px){.notebook{grid-template-columns:1fr}}"
    elif k == "editorial":
        body = (
            '<header class="campaign">'
            + title
            + '<div class="campaign-mark" aria-hidden="true">↺</div></header><div class="columns"><div><p class="lead">'
            + esc(d["intro"])
            + "</p>"
            + fig
            + "</div><div>"
            + stat
            + sec
            + "</div></div><h2>The Campaign in Public</h2>"
            + tab
        )
        css = ".campaign{position:relative;padding:35px 0 70px;border-bottom:4px solid var(--accent);margin-bottom:50px}.campaign h1{font-size:clamp(4rem,9vw,9rem);max-width:10ch;text-transform:uppercase}.campaign-mark{font-size:19rem;line-height:1;color:var(--accent);position:absolute;right:0;top:0;opacity:.2;z-index:-1}.campaign .lead{max-width:40ch}"
    elif k == "menu":
        body = (
            '<div class="menu"><aside class="wordmark">MORROW</aside><article>'
            + title
            + '<div class="courses">'
            + sec
            + "</div>"
            + intro
            + tab
            + "</article></div>"
        )
        css = '.menu{display:grid;grid-template-columns:180px 1fr;gap:70px}.wordmark{writing-mode:vertical-rl;transform:rotate(180deg);font:400 7rem/1 "Libre Baskerville";letter-spacing:.15em;color:var(--accent);border-left:1px solid var(--accent)}.courses{display:grid;grid-template-columns:1fr 1fr;gap:20px 50px}.courses h2{font:italic 400 1.5rem "Libre Baskerville";color:var(--accent)}.courses p{line-height:2.2}.menu h1{font-size:clamp(3rem,5vw,5rem)}@media(max-width:700px){.menu{display:block}.wordmark{writing-mode:horizontal-tb;transform:none;font-size:2rem;margin-bottom:36px;border:0}.courses{grid-template-columns:1fr}}'
    else:
        body = (
            '<header class="runbook">'
            + title
            + "<aside>"
            + stat
            + '</aside></header><div class="columns"><div>'
            + intro
            + sec
            + '</div><aside><div class="figure">'
            + motif(d)
            + "</div><h2>Proceed or Stop</h2>"
            + tab
            + '<div class="band"><h2>Stop Condition</h2><p>If any event is unmatched, pause delivery and preserve the checkpoint. Do not widen the replay batch.</p></div></aside></div>'
        )
        css = '.runbook{background:#102D38;color:white;padding:42px;display:grid;grid-template-columns:1fr 220px;gap:32px;margin:-64px -5vw 44px}.runbook aside{grid-row:1/4;grid-column:2;align-self:center;color:#A7EFF0}.runbook h1{font-size:clamp(2.5rem,4vw,4.4rem)}.reading h2{font-family:"Office Code Pro";font-size:1.3rem}@media(max-width:700px){.runbook{margin:-36px -22px 30px;display:block;padding:30px 22px}}'
    body += '<p class="closing">' + esc(d["closing"]) + "</p>"
    folder = OUT / "html" / d["id"]
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "styles.css").write_text(
        BASE + "\n" + css + color_css(d), encoding="utf8"
    )
    (folder / "tokens.css").write_text(
        ":root{--accent:"
        + d["accent"]
        + ";--paper:"
        + d["paper"]
        + ';--display:"'
        + d["font"]
        + '"}',
        encoding="utf8",
    )
    (OUT / "html" / f"{d['id']}.html").write_text(
        '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'
        + esc(d["title"])
        + '</title><link rel="stylesheet" href="../fonts/fonts.css"><link rel="stylesheet" href="'
        + d["id"]
        + '/tokens.css"><link rel="stylesheet" href="'
        + d["id"]
        + '/styles.css"><body><nav class="top"><a href="../index.html">Dazzler Design Gallery</a><button onclick="print()">Print This Example</button></nav><main class="canvas">'
        + body
        + "</main><footer>Fictional demonstration · Original Dazzler design · Apache-2.0. Local font licenses accompany the files. No information is transmitted.</footer></body></html>",
        encoding="utf8",
    )


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    x = OxmlElement("w:shd")
    x.set(qn("w:fill"), color.lstrip("#"))
    pr.append(x)


def para(
    parent,
    text,
    size=11,
    bold=False,
    color="17252B",
    font="Arial",
    after=8,
    italic=False,
    style=None,
):
    p = parent.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return p


def celltext(cell, text, size=11, bold=False, color="17252B", font="Arial"):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return p


def native_table(doc, d):
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    usable = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    for col, ratio in zip(t.columns, [0.25, 0.40, 0.35]):
        col.width = int(usable * ratio)
    margins = OxmlElement("w:tblCellMar")
    for edge in ["top", "left", "bottom", "right"]:
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    t._tbl.tblPr.append(margins)
    for i, h in enumerate(d["headers"]):
        shade(t.rows[0].cells[i], d["accent"])
        celltext(t.rows[0].cells[i], h, 10, True, "FFFFFF")
    for j, row in enumerate(d["rows"]):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            shade(cells[i], d["paper"] if j % 2 == 0 else "#FFFFFF")
            celltext(cells[i], v, 10, i == 0)
    for row in t.rows:
        x = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(x)
    x = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(x)
    return t


def word_document(d, destination=None):
    doc = Document()
    s = doc.sections[0]
    landscape = d["id"] in ["presentation", "family", "business"]
    s.page_width = Inches(11 if landscape else 8.5)
    s.page_height = Inches(8.5 if landscape else 11)
    s.top_margin = Inches(0.45 if d["id"] == "business" else 0.55)
    s.bottom_margin = Inches(0.45 if d["id"] == "business" else 0.55)
    s.left_margin = Inches(0.65)
    s.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    if d["id"] in ["fun", "school", "restaurant", "marketing", "legal"]:
        # A page-anchored vector surface prints without flattening editable copy.
        s.header.paragraphs[0].add_run()._r.append(
            parse_xml(
                '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
                '<v:rect id="DazzlerPageSurface" style="position:absolute;margin-left:0;margin-top:0;width:612pt;height:792pt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="'
                + d["paper"]
                + '" stroked="f"><v:fill color="'
                + d["paper"]
                + '"/></v:rect></w:pict>'
            )
        )
    para(doc, d["tag"], 9, True, d["accent"], after=14)
    k = d["layout"]
    if k == "poster":
        para(doc, "17 JUL", 76, True, d["accent"], d["word"], after=3)
        para(doc, heading(d["short"]), 36, True, "17252B", d["word"], style="Title")
        para(doc, d["lead"], 15, False, d["accent"], after=18)
        para(doc, d["label"], 12, True, d["accent"])
        para(doc, d["intro"], 10.5, after=8)
    elif k == "menu":
        p = para(doc, "MORROW", 43, False, d["accent"], "Georgia", after=7)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = para(
            doc,
            "Late Summer Supper",
            19,
            False,
            "17252B",
            "Georgia",
            italic=True,
            style="Title",
        )
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(doc, d["lead"], 11, False, "17252B", "Georgia", after=14)
        p = para(doc, d["stat"] + " · " + d["label"], 10, True, d["accent"], after=10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif k == "chronology":
        para(doc, d["title"], 30, False, "17252B", "Georgia", style="Title")
        para(doc, d["lead"], 14, False, d["accent"], "Georgia", italic=True, after=20)
    else:
        para(
            doc,
            d["title"],
            32 if not landscape else 34,
            True,
            d["accent"],
            d["word"],
            style="Title",
        )
        para(doc, d["lead"], 13, False, "17252B", d["word"], after=14)
    if k in ["decision", "proposal", "runbook"]:
        t = doc.add_table(rows=1, cols=2)
        t.columns[0].width = Inches(2.1)
        shade(t.cell(0, 0), d["accent"])
        celltext(t.cell(0, 0), d["stat"], 32, True, "FFFFFF", d["word"])
        celltext(t.cell(0, 0), d["label"], 10, False, "FFFFFF")
        celltext(t.cell(0, 1), d["intro"], 11)
        para(doc, "", 3, after=2)
    elif k not in ["menu", "poster"]:
        para(doc, d["intro"], 11, False, font=d["word"], after=12)
    if k == "planner":
        t = doc.add_table(rows=1, cols=3)
        for i, row in enumerate(d["rows"]):
            c = t.cell(0, i)
            shade(c, d["accent"] if i == 1 else d["paper"])
            ink = "FFFFFF" if i == 1 else "17252B"
            celltext(c, row[0], 24, True, ink)
            celltext(c, row[1], 13, True, ink)
            celltext(c, row[2], 11, False, ink)
        para(doc, "Pack and Check", 16, True, d["accent"], style="Heading 1")
        para(
            doc,
            "☐ Rain Jackets     ☐ Bottles     ☐ Medication     ☐ Chargers     ☐ Books",
            12,
        )
        for h, p in d["sections"][1:]:
            para(doc, heading(h), 13, True, d["accent"], style="Heading 1")
            para(doc, p, 10.5)
        x = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(x)
    elif k == "storyboard":
        t = doc.add_table(rows=1, cols=3)
        for i, (h, p) in enumerate(d["sections"]):
            c = t.cell(0, i)
            shade(c, d["accent"] if i == 1 else d["paper"])
            ink = "FFFFFF" if i == 1 else "17252B"
            celltext(c, f"0{i+1}", 40, True, ink, d["word"])
            celltext(c, heading(h), 15, True, ink, d["word"])
            celltext(c, p, 10.5, False, ink)
        para(doc, "", 3, after=3)
        native_table(doc, d)
    elif k == "menu":
        for h, p in d["sections"]:
            q = para(
                doc,
                heading(h),
                16,
                False,
                d["accent"],
                "Georgia",
                after=5,
                italic=True,
                style="Heading 1",
            )
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q = para(doc, p, 12, False, "17252B", "Georgia", after=17)
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        native_table(doc, d)
        para(doc, d["intro"], 9, after=6)
    elif k == "chronology":
        native_table(doc, d)
        for h, p in d["sections"]:
            para(doc, heading(h), 14, True, d["accent"], "Georgia", style="Heading 1")
            para(doc, p, 11, font="Georgia")
    elif k == "notebook":
        native_table(doc, d)
        for h, p in d["sections"]:
            para(doc, heading(h), 15, False, d["accent"], "Georgia", style="Heading 1")
            para(doc, p, 11, font="Georgia")
        para(
            doc,
            "Next Observation  ____________________________________",
            11,
            False,
            d["accent"],
        )
    elif k == "editorial":
        t = doc.add_table(rows=1, cols=1)
        shade(t.cell(0, 0), d["accent"])
        celltext(t.cell(0, 0), "KEEP THE GOOD THINGS", 27, True, "FFFFFF")
        celltext(t.cell(0, 0), "ASSESS  →  ESTIMATE  →  REPAIR", 14, False, "FFFFFF")
        for h, p in d["sections"]:
            para(doc, heading(h), 14, True, d["accent"], style="Heading 1")
            para(doc, p, 10.5)
        native_table(doc, d)
    elif k == "runbook":
        for h, p in d["sections"]:
            t = doc.add_table(rows=1, cols=2)
            t.columns[0].width = Inches(1.8)
            shade(t.cell(0, 0), d["accent"])
            celltext(t.cell(0, 0), heading(h), 13, True, "FFFFFF", "Consolas")
            shade(t.cell(0, 1), d["paper"])
            celltext(t.cell(0, 1), p, 10.5)
            para(doc, "", 2, after=3)
        native_table(doc, d)
        para(
            doc,
            "STOP: Any unmatched event means pause and preserve the checkpoint.",
            11,
            True,
            d["accent"],
        )
    else:
        for h, p in d["sections"]:
            para(
                doc,
                heading(h),
                13 if k == "proposal" else 14,
                True,
                d["accent"],
                d["word"],
                after=4 if k == "proposal" else 8,
                style="Heading 1",
            )
            para(doc, p, 10.5, after=5 if k == "proposal" else 8)
        native_table(doc, d)
    para(doc, d["closing"], 9, True, d["accent"], after=4)
    f = s.footer.paragraphs[0]
    f.text = "DAZZLER  /  FICTIONAL EXAMPLE  /  " + d["id"].upper()
    for r in f.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(d["accent"][1:])
    doc.core_properties.title = d["title"]
    doc.core_properties.author = "Jon Gosier"
    restyle_word(doc, d)
    doc.save(destination or OUT / "docx" / f"{d['id']}.docx")


def prepare():
    for p in ["html", "docx", "ui", "data", "previews"]:
        (OUT / p).mkdir(parents=True, exist_ok=True)
    if not (OUT / "fonts").exists():
        shutil.copytree(SKILL / "assets/templates/fonts", OUT / "fonts")
    (OUT / "index.html").write_text(
        "<title>Dazzler Review Staging</title><p>Internal rendering workspace.</p>",
        encoding="utf8",
    )
    old = STAGE / "previous-v0.27.0"
    old.mkdir(exist_ok=True)
    catalog = json.loads(
        (SKILL / "assets/templates/catalog.json").read_text(encoding="utf8")
    )
    for t in catalog["templates"]:
        target = old / (t["id"] + ".jpg")
        if not target.exists():
            target.write_bytes(
                subprocess.check_output(
                    [
                        "git",
                        "show",
                        "5defca0:skills/dazzler-frontend/assets/templates/previews/"
                        + t["id"]
                        + ".jpg",
                    ],
                    cwd=ROOT,
                )
            )
    return catalog


if __name__ == "__main__":
    prepare()
    for d in DOCS:
        document_html(d)
        word_document(d)
    (STAGE / "document-briefs.json").write_text(
        json.dumps(DOCS, indent=2), encoding="utf8"
    )
    print("Authored 10 Word and 10 HTML documents in staging")
