"""Author the 0.32 gallery from recorded briefs; rendering/review are separate gates.

Uses existing fictional source data and local demo controls, not a legacy layout builder.
"""

from pathlib import Path
import json, shutil, html, sys
from lxml import html as dom
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R = Path(__file__).resolve().parents[1]
L = R / "skills/dazzler-frontend/assets/templates"
S = R / "maintenance/gallery-production/v0.32.0"
B = R / "work/gallery-before-032"
H = html.escape
sys.path.insert(0, str(R / "skills/dazzler-frontend/scripts"))
from headings import heading

# Subject, composition, font, accent, paper, native orientation.
D = {
    "professional": (
        "The Release Gate",
        "A wide decision banner leads into a full-width reconciliation register; three ownership notes close the decision.",
        "Archivo",
        "#173D60",
        "#FFFFFF",
        "landscape",
    ),
    "legal": (
        "The Evidence Brief",
        "A portrait brief gives the conclusion and evidence register priority; dated records remain in one comparison table above the legal questions.",
        "Libre Baskerville",
        "#562B37",
        "#F6F0E5",
        "portrait",
    ),
    "business": (
        "The Scope Letter",
        "A restrained portrait offer pairs the fee with the proposal title; the phase register separates scope from commercial terms.",
        "Young Serif",
        "#22573D",
        "#FFFFFF",
        "portrait",
    ),
    "fun": (
        "The Evening Ticket",
        "A horizontal ticket places the event title beside the running order; a bright arrival band and three practical columns anchor the lower half.",
        "Oswald",
        "#6C2449",
        "#FFD85A",
        "landscape",
    ),
    "family": (
        "The Weekend Checklist",
        "A portrait itinerary puts packing and contingency notes before the day-by-day register so the reader can prepare before travelling.",
        "Archivo",
        "#264F87",
        "#E3EDF7",
        "portrait",
    ),
    "presentation": (
        "The Experiment Prospectus",
        "A portrait argument uses a narrow six-week marker, a central proposal, and a full-width measurement register before the decision conditions.",
        "Poppins",
        "#633A20",
        "#F7EBDD",
        "portrait",
    ),
    "school": (
        "The Trial Bench",
        "A wide laboratory sheet puts raw measurements first beside the research question; method and interpretation occupy a lower three-column band.",
        "Office Code Pro",
        "#713124",
        "#FFFFFF",
        "landscape",
    ),
    "marketing": (
        "The Campaign Dispatch",
        "A wide channel-first dispatch pairs the message with a tall capacity strip, then gives the launch channels the entire lower width.",
        "Archivo",
        "#223E76",
        "#D9E4FF",
        "landscape",
    ),
    "restaurant": (
        "The Seasonal Menu",
        "A narrow portrait menu alternates course headings and descriptions, followed by a quiet pairing register and service details.",
        "Libre Baskerville",
        "#385338",
        "#F5F2E7",
        "portrait",
    ),
    "technical": (
        "The Recovery Checkpoint",
        "A wide operations sheet puts the proceed-or-stop matrix before the three recovery steps; a red stop rule and batch size remain adjacent to the procedure.",
        "Office Code Pro",
        "#13545A",
        "#FFFFFF",
        "landscape",
    ),
}
U = {
    "webapp-workspace": (
        "The Manuscript Desk",
        "A wide manuscript leads; revision guidance moves into a labelled disclosure below the editor, with chapter context across the top.",
        "Libre Baskerville",
        "#303E67",
        "#F5F0E6",
    ),
    "webapp-board": (
        "The Production Register",
        "A dense vertical assignment register replaces the staircase; a right-hand release rail keeps deadline and completion beside the work.",
        "Archivo",
        "#4B3768",
        "#E9E1F6",
    ),
    "webapp-settings": (
        "Your Daily Digest",
        "A single-column settings form sits in a white sheet on a dark field; a compact selected-time strip replaces the tall clock and duplicated policy.",
        "Office Code Pro",
        "#254E54",
        "#17363C",
    ),
    "data-revenue": (
        "Six Years of Membership",
        "A wide chart leads a white analytical spread; measure controls sit above it and the exact values move beside the chart.",
        "Archivo",
        "#334BA0",
        "#FFFFFF",
    ),
    "data-operations": (
        "Station Readings",
        "A narrow station rail controls a dark instrument panel; a full-width exception strip closes the operational sequence.",
        "Office Code Pro",
        "#713B0F",
        "#EED6B3",
    ),
    "restaurant-fine-dining": (
        "Supper in the Garden",
        "An asymmetrical dark dining folio pairs a large botanical drawing with service information and an expandable menu on a pale sheet.",
        "Young Serif",
        "#254C36",
        "#173E2C",
    ),
    "restaurant-cafe": (
        "At the Coffee Counter",
        "A full-width receipt strip precedes a horizontal product shelf; the selected items and clear action stay beside the running subtotal.",
        "Oswald",
        "#63321C",
        "#F2CB9C",
    ),
    "restaurant-reservations": (
        "Choose Your Evening",
        "A form-first booking desk puts the request beside a vertical seating preference rail; access details remain directly below the zone choices.",
        "Poppins",
        "#58345C",
        "#F4EAF5",
    ),
    "restaurant-menu": (
        "Lunch at the Market",
        "A dark search rail accompanies a two-column dish catalog; the allergy instruction spans the page above the results.",
        "Archivo",
        "#38512B",
        "#FFFFFF",
    ),
    "business-portal": (
        "Review the Next Step",
        "The active document leads the review workspace; a compact milestone strip follows it, and review state belongs to each document separately.",
        "Poppins",
        "#4D386A",
        "#EEE9F6",
    ),
}


def save(p, value):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


BASE = """*{box-sizing:border-box}body{margin:0;background:var(--paper);color:#17212B;font:17px/1.6 "Work Sans",Arial,sans-serif}main{max-width:1300px;margin:auto;padding:45px 5vw 70px}h1,h2,h3{font-family:var(--display);line-height:1.15}h1{font-size:clamp(2.2rem,4.8vw,4.6rem);letter-spacing:-.025em;margin:0 0 22px;max-width:23ch}h2{font-size:1.4rem;margin:0 0 12px}p{max-width:60ch}a{color:inherit}.top,footer{padding:16px 5vw;font-size:13px}.top{display:flex;justify-content:space-between;gap:16px;border-bottom:1px solid #79858B}.eyebrow{font-size:12px;letter-spacing:.12em;font-weight:700;text-transform:uppercase}.lead{font-size:1.2rem;max-width:55ch}.pair{display:grid;grid-template-columns:1fr 1fr;gap:35px;align-items:start}.thirds{display:grid;grid-template-columns:repeat(3,1fr);gap:28px}.band{margin:28px 0;padding:24px;background:var(--accent);color:white}.sheet{background:#fff;padding:30px}.note{border-top:2px solid var(--accent);padding-top:18px;margin:22px 0}.note p{white-space:pre-line}.metric{padding:20px 0}.metric strong{display:block;font:clamp(2.2rem,4vw,4rem)/1.1 var(--display);font-variant-numeric:tabular-nums}.metric p{margin:8px 0}.closing{margin-top:30px;padding-top:18px;border-top:1px solid #79858B;font-size:14px}.table-wrap{overflow:auto}table{border-collapse:collapse;width:100%;font-size:15px;text-align:left}th,td{padding:12px 10px;border-bottom:1px solid #899398;vertical-align:top}thead{background:var(--accent);color:#fff}button,input,select,textarea{font:inherit}button{background:var(--accent);color:white;border:2px solid var(--accent);padding:12px 18px;cursor:pointer}button[aria-pressed=false]{background:white;color:var(--accent)}button:disabled{opacity:.65}label{display:block;font-weight:600}input:not([type=checkbox]),select{width:100%;padding:12px;background:white;color:#17212B;border:1px solid #65747A;margin:8px 0 20px}input[type=checkbox]{width:20px;height:20px;vertical-align:middle}.actions{display:flex;flex-wrap:wrap;gap:12px;margin:20px 0}textarea{width:100%;resize:vertical}figure{margin:0;min-width:0}svg{width:100%;height:auto}[role=status]{min-height:1.6em}:focus-visible{outline:3px solid #B03B00;outline-offset:4px}small{font-size:14px}[hidden]{display:none!important}summary{cursor:pointer;font-weight:600;padding:16px 0}section,aside,article,div{min-width:0}@media(max-width:750px){main{padding:28px 20px}.pair,.thirds{display:block}.sheet{padding:22px}.top{flex-wrap:wrap}h1{font-size:2.4rem}th,td{padding:10px 8px}}@media print{.top{display:none}main{padding:0}h1{font-size:30pt}footer{font-size:9pt}}"""


def table(d):
    return (
        '<div class="table-wrap" tabindex="0" role="region" aria-label="Exact Data"><table><thead><tr>'
        + "".join('<th scope="col">' + H(v) + "</th>" for v in d["headers"])
        + "</tr></thead><tbody>"
        + "".join(
            "<tr>"
            + "".join(
                ('<th scope="row">' if i == 0 else "<td>")
                + H(v)
                + ("</th>" if i == 0 else "</td>")
                for i, v in enumerate(row)
            )
            + "</tr>"
            for row in d["rows"]
        )
        + "</tbody></table></div>"
    )


def web_doc(d):
    k = d["id"]
    title = (
        '<header><p class="eyebrow">'
        + H(d["tag"])
        + "</p><h1>"
        + H(d["title"])
        + '</h1><p class="lead">'
        + H(d["lead"])
        + "</p></header>"
    )
    intro = "<p>" + H(d["intro"]) + "</p>"
    metric = (
        '<aside class="metric"><strong>'
        + H(d["stat"])
        + "</strong><p>"
        + H(d["label"])
        + "</p></aside>"
    )
    notes = [
        '<section class="note"><h2>'
        + H(heading(a))
        + "</h2><p>"
        + H(b)
        + "</p></section>"
        for a, b in d["sections"]
    ]
    ledger = table(d)
    css = ""
    if k == "professional":
        body = (
            title
            + '<div class="band pair">'
            + metric
            + intro
            + "</div>"
            + ledger
            + '<div class="thirds">'
            + "".join(notes)
            + "</div>"
        )
    elif k == "legal":
        body = (
            '<div class="brief">' + title + intro + ledger + "".join(notes) + "</div>"
        )
        css = ".brief{max-width:850px;margin:auto}.brief h1{font-size:3.8rem}.brief .note{margin-left:12%;border-top:0;border-left:3px solid var(--accent);padding-left:22px}"
    elif k == "business":
        body = (
            '<div class="letter">'
            + title
            + '<div class="fee">'
            + metric
            + notes[2]
            + "</div>"
            + intro
            + ledger
            + notes[0]
            + notes[1]
            + "</div>"
        )
        css = ".letter{max-width:900px;margin:auto}.fee{display:grid;grid-template-columns:1fr 2fr;gap:35px;border-block:2px solid var(--accent);margin:30px 0}.fee .note{border:0}"
    elif k == "fun":
        body = (
            '<div class="ticket"><div>'
            + title
            + intro
            + '</div><section class="sheet"><h2>The Running Order</h2>'
            + ledger
            + '</section></div><div class="band">'
            + metric
            + '</div><div class="thirds">'
            + "".join(notes)
            + "</div>"
        )
        css = ".ticket{display:grid;grid-template-columns:1.2fr 1fr;gap:45px}.ticket h1{font-size:clamp(3rem,6vw,6rem)}.band .metric{display:flex;align-items:center;gap:30px}"
    elif k == "family":
        body = (
            title
            + '<div class="packing"><div>'
            + notes[0]
            + notes[1]
            + notes[2]
            + '</div><section class="sheet"><h2>The Weekend Plan</h2>'
            + ledger
            + intro
            + "</section></div>"
        )
        css = ".packing{display:grid;grid-template-columns:1fr 1.5fr;gap:45px}.packing .note:first-child{border-top:10px solid var(--accent)}"
    elif k == "presentation":
        body = (
            '<div class="prospectus"><aside>'
            + metric
            + "</aside><div>"
            + title
            + intro
            + notes[0]
            + "</div></div>"
            + ledger
            + '<div class="pair">'
            + notes[1]
            + notes[2]
            + "</div>"
        )
        css = ".prospectus{display:grid;grid-template-columns:1fr 3fr;gap:45px}.prospectus aside{border-top:12px solid var(--accent)}"
    elif k == "school":
        marks = "".join(
            f'<text x="0" y="{30+i*48}">{H(r[0])}</text><line x1="160" x2="{160+int(r[2].split()[0])}" y1="{25+i*48}" y2="{25+i*48}" stroke="currentColor"/><circle cx="{160+int(r[2].split()[0])}" cy="{25+i*48}" r="5"/><text x="{172+int(r[2].split()[0])}" y="{30+i*48}">{H(r[2])}</text>'
            for i, r in enumerate(d["rows"])
        )
        chart = (
            '<figure class="drainage"><svg role="img" aria-labelledby="drainage-title drainage-desc" viewBox="0 0 420 170"><title id="drainage-title">Median Drainage Times</title><desc id="drainage-desc">Sand 40 seconds; garden soil 87 seconds; clay-rich soil 168 seconds. Dots share a zero baseline.</desc><g font-family="Work Sans,Arial" font-size="15" fill="currentColor">'
            + marks
            + '<text x="155" y="162">0</text><text x="310" y="162">Seconds</text></g></svg><figcaption>Median time to collect 100 mL. Exact trials above.</figcaption></figure>'
        )
        body = (
            '<div class="lab"><section class="sheet"><p class="eyebrow">MEASUREMENTS / SECONDS</p>'
            + ledger
            + chart
            + metric
            + "</section><div>"
            + title
            + intro
            + '</div></div><div class="thirds">'
            + "".join(notes)
            + "</div>"
        )
        css = ".lab{display:grid;grid-template-columns:1.2fr 1fr;gap:45px}.lab .sheet{border:2px solid var(--accent)}.lab h1{font-size:3.2rem}.drainage{margin-top:22px}.drainage figcaption{font-size:14px}"
    elif k == "marketing":
        body = (
            '<div class="dispatch"><div>'
            + title
            + intro
            + '</div><aside class="band">'
            + metric
            + notes[0]
            + "</aside></div>"
            + ledger
            + '<div class="pair">'
            + notes[1]
            + notes[2]
            + "</div>"
        )
        css = ".dispatch{display:grid;grid-template-columns:2fr 1fr;gap:45px}.dispatch .band{margin:0 0 30px}.dispatch .note{border-color:white}"
    elif k == "restaurant":
        body = (
            '<div class="menu">'
            + title
            + "".join(notes)
            + ledger
            + metric
            + intro
            + "</div>"
        )
        css = ".menu{max-width:730px;margin:auto;text-align:center}.menu p{margin-left:auto;margin-right:auto}.menu .note{border:0;padding:18px 0}.menu h1{font-size:3.6rem}.menu h2{color:var(--accent)}.menu table{text-align:left}"
    else:
        body = (
            '<div class="pair">'
            + title
            + '<aside class="stop"><h2>Stop on Any Failed Check</h2><p>Pause again and retain the evidence if the batch fails.</p>'
            + metric
            + "</aside></div>"
            + ledger
            + '<div class="thirds">'
            + "".join(notes)
            + "</div>"
            + intro
        )
        css = ".stop{border-left:8px solid #A22628;background:#FFF0EE;padding:24px}.stop h2{color:#8A2024}"
    css += "@media(max-width:750px){.ticket,.packing,.prospectus,.lab,.dispatch,.fee{display:block}.brief h1,.menu h1,.lab h1{font-size:2.4rem}.brief .note{margin-left:0}.band .metric{display:block}.lab .sheet{margin-bottom:25px}}"
    return body + '<p class="closing">' + H(d["closing"]) + "</p>", css


def native(d, config):
    _, _, _, accent, paper, orientation = config
    doc = Document()
    sec = doc.sections[0]
    landscape = orientation == "landscape"
    for element in list(doc.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    sec.page_width = Inches(11.7 if landscape else 8.3)
    sec.page_height = Inches(8.3 if landscape else 11.7)
    sec.top_margin = sec.bottom_margin = Inches(0.48)
    sec.left_margin = sec.right_margin = Inches(0.55)
    face = (
        "Consolas"
        if d["id"] in ("technical", "school")
        else ("Georgia" if d["id"] in ("legal", "restaurant", "business") else "Arial")
    )
    for st in ["Normal", "Title", "Heading 1"]:
        doc.styles[st].font.name = face
        doc.styles[st].font.size = Pt(10)
        doc.styles[st].font.color.rgb = RGBColor.from_string("17212B")
        for e in list(doc.styles[st].element.iter(qn("w:pBdr"))):
            e.getparent().remove(e)

    def p(where, text, size=11, bold=False, color="17212B", style=None):
        el = where.add_paragraph(style=style)
        el.paragraph_format.space_after = Pt(7)
        el.paragraph_format.line_spacing = 1.08
        if len(text) > 80 and size <= 12 and not bold:
            width = getattr(where, "width", None) or (
                sec.page_width - sec.left_margin - sec.right_margin
            )
            el.paragraph_format.right_indent = Pt(max(0, width / 12700 - size * 31))
        run = el.add_run(text)
        run.bold = bold
        run.font.name = face
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        return el

    def shade(cell, color):
        e = OxmlElement("w:shd")
        e.set(qn("w:fill"), color.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(e)

    def grid(widths):
        t = doc.add_table(rows=1, cols=len(widths))
        t.autofit = False
        for col, c, w in zip(t.columns, t.rows[0].cells, widths):
            col.width = Inches(w)
            c.width = Inches(w)
        return t.rows[0].cells

    def title(w):
        p(w, d["title"], 27 if landscape else 29, True, style="Title")
        p(w, d["lead"], 11, True, accent)

    def note(w, n):
        p(w, heading(n[0]), 13, True, accent, "Heading 1")
        p(w, n[1])

    def metric(w):
        p(w, d["stat"], 25, True, accent)
        p(w, d["label"])

    def ledger(w):
        t = w.add_table(rows=1, cols=len(d["headers"]))
        t.autofit = True
        for c, v in zip(t.rows[0].cells, d["headers"]):
            shade(c, accent)
            p(c, v, 10, True, "FFFFFF")
        e = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(e)
        for row in d["rows"]:
            for c, v in zip(t.add_row().cells, row):
                p(c, v, 10)
        for row in t.rows:
            e = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(e)

    k = d["id"]
    p(doc, d["tag"], 9, True, accent)
    if k == "professional":
        title(doc)
        cells = grid([2.2, 8.2])
        metric(cells[0])
        p(cells[1], d["intro"])
        ledger(doc)
        for c, n in zip(grid([3.45] * 3), d["sections"]):
            note(c, n)
    elif k == "legal":
        title(doc)
        p(doc, d["intro"])
        ledger(doc)
        for n in d["sections"]:
            note(doc, n)
    elif k == "business":
        title(doc)
        metric(doc)
        p(doc, d["intro"])
        ledger(doc)
        for n in d["sections"]:
            note(doc, n)
    elif k == "fun":
        left, right = grid([5.2, 5.2])
        title(left)
        p(left, d["intro"])
        ledger(right)
        metric(doc)
        for c, n in zip(grid([3.45] * 3), d["sections"]):
            note(c, n)
    elif k == "family":
        title(doc)
        for n in d["sections"]:
            note(doc, n)
        ledger(doc)
        p(doc, d["intro"])
    elif k == "presentation":
        metric(doc)
        title(doc)
        p(doc, d["intro"])
        note(doc, d["sections"][0])
        ledger(doc)
        for n in d["sections"][1:]:
            note(doc, n)
    elif k == "school":
        left, right = grid([5.2, 5.2])
        ledger(left)
        metric(left)
        title(right)
        p(right, d["intro"])
        for c, n in zip(grid([3.45] * 3), d["sections"]):
            note(c, n)
    elif k == "marketing":
        left, right = grid([7.2, 3.2])
        title(left)
        p(left, d["intro"])
        metric(right)
        note(right, d["sections"][0])
        ledger(doc)
        for c, n in zip(grid([5.2] * 2), d["sections"][1:]):
            note(c, n)
    elif k == "restaurant":
        title(doc)
        for n in d["sections"]:
            note(doc, n)
        ledger(doc)
        metric(doc)
        p(doc, d["intro"])
    else:
        left, right = grid([7.2, 3.2])
        title(left)
        p(right, "Stop on Any Failed Check", 13, True, "A22628")
        metric(right)
        ledger(doc)
        for c, n in zip(grid([3.45] * 3), d["sections"]):
            note(c, n)
        p(doc, d["intro"])
    p(doc, d["closing"], 9, True, accent)
    p(
        sec.footer,
        "Fictional Dazzler example / Apache-2.0 / Editable text and tables",
        8,
    )
    # A page-anchored shape survives PDF export and printing, unlike Word page color.
    from lxml import etree

    sec.header.paragraphs[0].add_run()._r.append(
        etree.fromstring(
            f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:rect id="DazzlerPageSurface" style="position:absolute;margin-left:0;margin-top:0;width:{sec.page_width/12700}pt;height:{sec.page_height/12700}pt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="{paper}" stroked="f"><v:fill color="{paper}"/></v:rect></w:pict>'
        )
    )
    doc.save(L / "docx" / f"{k}.docx")
    return face


def page(title, body, css, font, accent, paper, depth, scripts=""):
    markup = (
        f'<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{H(title)}</title><link rel="stylesheet" href="{depth}fonts/fonts.css"><style>:root{{--display:"{font}";--accent:{accent};--paper:{paper}}}'
        + BASE
        + css
        + '</style><body><nav class="top"><a href="'
        + depth
        + 'index.html">Dazzler Design Gallery</a><span>Fictional Example</span></nav><main>'
        + body
        + "</main><footer>Original Dazzler design · Apache-2.0 · Local font notices included.</footer>"
        + scripts
        + "</body></html>"
    )
    tree = dom.fromstring(markup)
    for el in tree.xpath("//h1|//h2|//h3"):
        if not len(el):
            el.text = heading(el.text or "")
    return "<!doctype html>" + dom.tostring(tree, encoding="unicode")


def ui(k, c):
    title, why, font, accent, paper = c
    old = dom.parse(
        str(B / "ui" / k / "index.html"), parser=dom.HTMLParser(encoding="utf-8")
    )

    def get(x):
        return "".join(
            dom.tostring(e, encoding="unicode", with_tail=False) for e in old.xpath(x)
        )

    def ident(x):
        return get('//*[@id="' + x + '"]')

    def cls(x):
        return get(
            '//*[contains(concat(" ",normalize-space(@class)," ")," ' + x + ' ")]'
        )

    lead = old.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," lead ")]')[
        0
    ].text_content()
    head = (
        '<header><p class="eyebrow">'
        + H(k.replace("-", " / "))
        + "</p><h1>"
        + H(title)
        + '</h1><p class="lead">'
        + H(lead)
        + "</p></header>"
    )
    css = ""
    if k == "webapp-workspace":
        body = (
            cls("chapterline")
            + head
            + cls("writing")
            + '<details class="revision-drawer"><summary>Revision Notes</summary>'
            + cls("revision")
            + "</details>"
        )
        css = '.chapterline{display:flex;gap:30px;border-bottom:1px solid;padding-bottom:18px;margin-bottom:35px}.writing{max-width:950px;background:white;padding:35px;margin:auto}.writing textarea{min-height:390px;padding:22px;border:1px solid #8B8F96;font:18px/1.8 "Libre Baskerville"}.revision-drawer{max-width:950px;margin:25px auto;border-top:2px solid var(--accent)}'
    elif k == "webapp-board":
        body = (
            head
            + '<div class="production"><section>'
            + ident("tasks")
            + ident("status")
            + '</section><aside class="release">'
            + cls("deadline")
            + ident("done")
            + cls("release-note")
            + "</aside></div>"
        )
        css = ".production{display:grid;grid-template-columns:3fr 1fr;gap:40px}.task{display:grid;grid-template-columns:2fr 1fr auto;gap:15px;align-items:center;border-top:1px solid;padding:22px 0}.task h2{font-size:1.25rem}.release{background:white;padding:25px;border-top:9px solid var(--accent)}.deadline{font-size:1.8rem;font-weight:bold}"
    elif k == "webapp-settings":
        body = (
            '<div class="settings-sheet">'
            + head
            + '<div class="selected"><span>Selected Digest Time</span>'
            + ident("time")
            + "</div>"
            + ident("settings")
            + "</div>"
        )
        css = "body{color:white}.settings-sheet{max-width:760px;margin:auto;background:white;color:#17212B;padding:38px}.selected{display:flex;justify-content:space-between;gap:20px;background:#E5F1EF;padding:18px;margin:25px 0}.selected strong{font-size:1.6rem}fieldset{padding:20px;margin:25px 0;border:1px solid #738184}fieldset label{margin:14px 0}"
    elif k == "data-revenue":
        body = (
            head
            + cls("actions")
            + '<div class="analytics"><figure><h2 id="chart-title"></h2>'
            + ident("chart")
            + ident("status")
            + "</figure><section><h2>Annual Values</h2>"
            + ident("values")
            + "</section></div><p>Synthetic observations. Measures use separate scales and do not establish causality.</p>"
        )
        css = ".analytics{display:grid;grid-template-columns:1.6fr 1fr;gap:45px;border-block:2px solid var(--accent);padding:30px 0}.bars{display:flex;align-items:end;gap:12px;height:340px}.bar{display:flex;flex:1;flex-direction:column;align-items:center;font-size:13px;gap:8px}.bar i{display:block;background:var(--accent);width:100%}figcaption{font-weight:bold;font-size:15px}"
    elif k == "data-operations":
        body = (
            head
            + '<div class="instrument">'
            + cls("stations")
            + ident("inspector")
            + "</div>"
            + cls("exception")
        )
        css = ".instrument{display:grid;grid-template-columns:1fr 3fr;gap:24px}.stations{display:flex;flex-direction:column;gap:12px}#inspector{background:#222B31;color:white;padding:40px}.reading-value{font-size:5rem;font-variant-numeric:tabular-nums}.exception{margin-top:28px;border-left:8px solid #A5222B;background:#FFF0EE;padding:25px}"
    elif k == "restaurant-fine-dining":
        body = (
            '<div class="dining"><aside>'
            + cls("botanical")
            + cls("service")
            + '</aside><section class="sheet">'
            + head
            + get("//main//header/p[not(@class)]")
            + ident("courses-button")
            + ident("courses")
            + "</section></div>"
        )
        css = "body{color:white}.dining{display:grid;grid-template-columns:1fr 1.4fr;gap:60px}.dining .sheet{color:#17212B;padding:45px}.botanical{max-width:360px;margin:20px auto}.botanical path{stroke:#E6DBA8}.botanical g{fill:#E6DBA8}.service{border-top:1px solid;padding-top:25px}.courses{list-style:none;padding:0}.courses li{border-bottom:1px solid;padding:20px 0}.courses strong{display:block;font-size:1.3rem}"
    elif k == "restaurant-cafe":
        body = (
            head
            + '<aside class="receipt"><div><h2>Your Counter Order</h2>'
            + ident("cart")
            + '</div><div><p>Subtotal Before Tax</p><strong id="total">$0.00</strong>'
            + ident("clear")
            + "</div><p>Local demo. No order or payment is submitted.</p></aside>"
            + cls("actions")
            + ident("status")
            + ident("products")
        )
        css = '.receipt{display:grid;grid-template-columns:2fr 1fr;gap:22px;border-block:3px solid var(--accent);padding:20px 0}.receipt>p{grid-column:1/-1;margin:0}#total{font:3rem "Oswald";display:block;margin-bottom:15px}#products{display:grid;grid-template-columns:repeat(3,1fr);gap:26px}#products article{background:#FFF9F0;padding:25px;border-top:8px solid var(--accent)}'
    elif k == "restaurant-reservations":
        body = (
            head
            + '<div class="reservation"><div class="sheet">'
            + ident("booking")
            + "</div><aside><h2>Seating Preference</h2>"
            + cls("zones")
            + cls("access")
            + "</aside></div>"
        )
        css = ".reservation{display:grid;grid-template-columns:1.4fr 1fr;gap:45px}.zones{display:flex;flex-direction:column;gap:15px}.zones button{display:flex;justify-content:space-between;align-items:center;gap:15px;min-height:85px}.zone-mark{font-size:1.5rem}.access{margin-top:25px;border-top:2px solid;padding-top:20px}.requestline{display:block}"
    elif k == "restaurant-menu":
        controls = get('//*[@id="search"]/..')
        # Preserve only the search label/input and dietary controls, not the former toolbar wrapper.
        controls = get('//label[@for="search"]') + get(
            '//label[.//input[@type="checkbox"]]'
        )
        body = (
            head
            + cls("allergy")
            + '<div class="market"><aside class="search-rail">'
            + controls
            + ident("result-count")
            + "</aside><section><h2>Today’s Dishes</h2>"
            + ident("menu-items")
            + "</section></div>"
        )
        css = ".allergy{border-left:6px solid #9C2934;padding-left:20px;margin:25px 0}.market{display:grid;grid-template-columns:1fr 3fr;gap:35px}.search-rail{background:#253629;color:white;padding:25px;align-self:start}.search-rail label{margin-bottom:12px}#menu-items{display:grid;grid-template-columns:1fr 1fr;gap:25px}#menu-items article{border-top:2px solid var(--accent);padding:20px 0}#menu-items h2{font-size:1.45rem}"
    else:
        body = (
            head
            + '<div class="review"><aside>'
            + cls("documents")
            + '</aside><section class="sheet">'
            + get('//*[contains(@class,"review-sheet")]/article')
            + "</section></div>"
            + cls("milestone-ribbon")
            + "<p>No approval or message is sent. Review marks apply only to this preview.</p>"
        )
        css = ".review{display:grid;grid-template-columns:1fr 3fr;gap:30px}.documents{display:flex;flex-direction:column;gap:12px}.milestone-ribbon{margin-top:30px}.milestone-ribbon ol{display:flex;justify-content:space-between;gap:20px;list-style:none;padding:25px 0;border-top:2px solid var(--accent)}#preview{font-size:1.2rem;line-height:1.8}.review .sheet{border-left:6px solid var(--accent)}"
    css += "@media(max-width:750px){.production,.analytics,.instrument,.dining,.reservation,.market,.review,.receipt{display:block}.task{display:block}.task span{display:block;margin:8px 0}.task button{display:block;margin-top:12px}.botanical{max-width:150px}.release,.exception{margin-top:25px}.chapterline{display:block}.chapterline span{display:block}.settings-sheet{padding:22px}.selected{display:block}.selected strong{display:block}.stations{margin-bottom:20px}.dining .sheet{padding:22px}#products,#menu-items{grid-template-columns:1fr}.bars{gap:6px}.bar{font-size:11px}.receipt button{margin:15px 0}.search-rail{margin-bottom:25px}.documents{flex-direction:row;flex-wrap:wrap;margin-bottom:20px}.milestone-ribbon ol{display:block}.milestone-ribbon li{margin:15px 0}.writing{padding:20px}.writing textarea{padding:12px}.reading-value{font-size:3.5rem}}"
    scripts = get("//script")
    if k == "business-portal":
        # A review belongs to the selected document; switching tabs cannot carry approval across it.
        start = scripts.index('      document.querySelectorAll("[data-doc]")')
        end = scripts.index("</script>", start)
        scripts = scripts[:start] + """let active = "Materials";
      const reviewed = new Set();
      const syncReview = () => {
        document.querySelector("#approve").textContent = reviewed.has(active) ? "Undo Local Review" : "Mark Reviewed Locally";
        document.querySelector("#status").textContent = reviewed.has(active) ? active + " marked reviewed in this preview only." : active + " awaits local review.";
        document.querySelectorAll("[data-doc]").forEach(b => b.setAttribute("aria-pressed", String(b.dataset.doc === active)));
      };
      document.querySelectorAll("[data-doc]").forEach(b => b.onclick = () => {
        active = b.dataset.doc;
        document.querySelector("#preview-title").textContent = active;
        document.querySelector("#preview").textContent = docs[active];
        syncReview();
      });
      document.querySelector("#approve").onclick = () => {
        if (reviewed.has(active)) reviewed.delete(active); else reviewed.add(active);
        syncReview();
      };
      syncReview();
    """ + scripts[end:]
    prompt = f"Design a fictional {k} example for this task: {lead} {why} Preserve the supplied data, keyboard controls, recovery actions and local-only behavior. Use {font} and a {paper} page with {accent} accents."
    data = json.loads((B / "ui" / k / "template.json").read_text(encoding="utf-8"))
    data.update(title=title, prompt=prompt)
    save(L / "ui" / k / "template.json", data)
    tree = dom.fromstring(
        page(title, body, css, font, accent, paper, "../../", scripts)
    )
    style = tree.xpath("//style")[0]
    (L / "ui" / k / "styles.css").write_text(style.text, encoding="utf-8")
    style.tag = "link"
    style.text = None
    style.set("rel", "stylesheet")
    style.set("href", "styles.css")
    tree.xpath('//*[@id="template-data"]')[0].text = json.dumps(
        data, ensure_ascii=False
    ).replace("<", "\\u003c")
    (L / "ui" / k / "index.html").write_text(
        "<!doctype html>" + dom.tostring(tree, encoding="unicode"), encoding="utf-8"
    )
    return prompt


def main():
    S.mkdir(parents=True, exist_ok=True)
    if not B.exists():
        shutil.copytree(L, B)
    prev = S / "previous-v0.31.1"
    prev.mkdir(exist_ok=True)
    for p in (B / "previews").glob("*.jpg"):
        shutil.copy2(p, prev / p.name)
    briefs = []
    for k, c in D.items():
        d = json.loads((B / "data" / f"{k}.json").read_text(encoding="utf-8"))
        title, why, font, accent, paper, orientation = c
        prompt = f'Create a fictional {k} document and HTML edition: {d["lead"]} {why} Preserve every source number and qualification. Use {font} for the web; keep the Word text and tables editable and inspect every rendered page.'
        d.update(font=font, accent=accent, paper=paper, layout=title, prompt=prompt)
        d["design"].update(
            accent=accent,
            secondary=accent,
            paper=paper,
            text="#17212B",
            accentText=accent,
            onAccent="#FFFFFF",
        )
        d["fontSelection"].update(id=font.lower().replace(" ", "-"), css_family=font)
        body, css = web_doc(d)
        (L / "html" / f"{k}.html").write_text(
            page(d["title"], body, css, font, accent, paper, "../"), encoding="utf-8"
        )
        nativeface = native(d, c)
        d["word"] = nativeface
        save(L / "data" / f"{k}.json", d)
        briefs.append(
            dict(
                id=k,
                prompt=prompt,
                direction=title,
                rationale=why,
                font=font,
                word=nativeface,
                orientation=orientation,
                surface=paper,
                accent=accent,
                alternatives=[
                    "Retain the prior composition",
                    title,
                    "Linear plain-text reference",
                ],
                selectionReason=why,
                structuralChanges=[
                    why,
                    "Changes page orientation and reading sequence from the previous gallery.",
                    "Moves the evidence register relative to the title and supporting notes.",
                ],
            )
        )
    save(S / "document-briefs.json", briefs)
    briefs = []
    for k, c in U.items():
        prompt = ui(k, c)
        title, why, font, accent, paper = c
        briefs.append(
            dict(
                id=k,
                prompt=prompt,
                direction=title,
                rationale=why,
                font=font,
                surface=paper,
                accent=accent,
                alternatives=[
                    "Retain the prior composition",
                    title,
                    "Single-column linear interface",
                ],
                selectionReason=why,
                structuralChanges=[
                    why,
                    "Regroups the primary control and its feedback into a new reading sequence.",
                    "Changes the balance of content, navigation and supporting information.",
                ],
            )
        )
    save(S / "interface-briefs.json", briefs)
    print("Authored 10 Word, 10 HTML and 10 UI examples; rendering and review pending.")


if __name__ == "__main__":
    main()
