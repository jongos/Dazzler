"""Author release 0.30 examples; capture and visual review are separate steps."""

from pathlib import Path
import html
import json
import shutil
import sys
from lxml import html as dom
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R = Path(__file__).resolve().parents[1]
S = R / "maintenance/gallery-production/v0.30.0"
O = S / "artifacts"
OLD = R / "work/gallery-before-030"
H = html.escape
sys.path.insert(0, str(R / "skills/dazzler-frontend/scripts"))
from headings import heading


def titled(page):
    tree = dom.fromstring(page)
    for element in tree.xpath("//h1|//h2|//h3"):
        if not len(element):
            element.text = heading(element.text or "")
    return "<!doctype html>" + dom.tostring(tree, encoding="unicode")


# Each direction is authored around the existing fictional reader task.
# Tuple: concept, web face, native face, accent, page, geometry, organizing idea.
D = {
    "professional": (
        "The Exception Dossier",
        "Office Code Pro",
        "Consolas",
        "#A52620",
        "#FFFFFF",
        "portrait",
        "Separate the 100 unresolved records in a narrow evidence column; lead with the explicit hold before the release register.",
    ),
    "legal": (
        "A Chronology for Counsel",
        "Archivo",
        "Arial",
        "#183A71",
        "#FFFFFF",
        "landscape",
        "Read three dated records left to right, then distinguish the available record from the actions needed below.",
    ),
    "business": (
        "Scope Across Three Stages",
        "Poppins",
        "Arial",
        "#4D2298",
        "#EAC7FF",
        "landscape",
        "A proposal foldout separates fee and boundaries from the three staged deliverables; the fee is a side annotation, not the page opener.",
    ),
    "fun": (
        "The Neighborhood Poster",
        "Archivo",
        "Arial",
        "#452378",
        "#FF9068",
        "portrait",
        "A vertical event poster makes arrival details the top band and the program a stepped running order under a giant title.",
    ),
    "family": (
        "The Weekend Route",
        "Young Serif",
        "Georgia",
        "#145A3A",
        "#FFF1AF",
        "landscape",
        "A horizontal route joins the three days; packing and change instructions sit in separate bottom pockets.",
    ),
    "presentation": (
        "Experiment and Decision",
        "Libre Baskerville",
        "Georgia",
        "#173D75",
        "#C2E9FF",
        "landscape",
        "A two-panel argument pairs the proposed experiment with the decision criteria; supporting measures sit across the lower edge.",
    ),
    "school": (
        "Drainage Field Notes",
        "Young Serif",
        "Georgia",
        "#9A2D0B",
        "#FFF5E4",
        "portrait",
        "A ruled field notebook starts with the experimental question, then nests the exact results between observations and limitations.",
    ),
    "marketing": (
        "The Repair Broadside",
        "Oswald",
        "Arial",
        "#1551A4",
        "#FFFFFF",
        "portrait",
        "An oversized campaign statement leads into a narrow numbered channel register and an indented offer story.",
    ),
    "restaurant": (
        "The Supper Ledger",
        "Office Code Pro",
        "Consolas",
        "#75243A",
        "#F7D2DD",
        "landscape",
        "Three course columns align dishes across a wide supper ledger; service time occupies a small running header and pairings close the page.",
    ),
    "technical": (
        "Recovery Control Sheet",
        "Archivo",
        "Arial",
        "#075866",
        "#D3FFF0",
        "portrait",
        "A prominent stop strip precedes a sequential procedure with a separate right-hand verification register; evidence stays adjacent to the decision.",
    ),
}


def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


BASE = """*{box-sizing:border-box}body{margin:0;background:var(--paper);color:#17212B;font:17px/1.6 "Work Sans",Arial}main{max-width:1320px;padding:50px 5vw 65px;margin:auto}h1,h2,h3{font-family:var(--display);line-height:1.12}h1{font-size:clamp(2.5rem,5.8vw,5.8rem);letter-spacing:-.035em;margin:0 0 28px;max-width:20ch}h2{font-size:1.6rem;margin:0 0 14px}p{max-width:60ch}a{color:inherit}.top,footer{padding:18px 5vw;font-size:13px;border-bottom:1px solid currentColor}.top{display:flex;justify-content:space-between;gap:20px}.eyebrow{font-size:12px;font-weight:700;letter-spacing:.12em;text-transform:uppercase}.lead{font-size:1.3rem;max-width:50ch}.spread{display:grid;grid-template-columns:1fr 1fr;gap:36px}.wide{grid-column:1/-1}.ink{background:var(--accent);color:white;padding:30px}.sheet{background:white;padding:28px}.notes{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:28px;margin:30px 0}.note{padding-top:18px;border-top:2px solid var(--accent)}.note p{white-space:pre-line}.stat{font-family:var(--display);font-size:clamp(2.3rem,5vw,4.8rem);line-height:1.1}.table-wrap,.table-scroll{overflow:auto}table{border-collapse:collapse;width:100%;font-size:15px;text-align:left}th,td{padding:14px 12px;border-bottom:1px solid #818B94;vertical-align:top}thead{background:var(--accent);color:white}.closing{margin-top:30px;border-top:3px solid var(--accent);padding-top:18px;font-weight:600}button,input,select,textarea{font:inherit}button{background:var(--accent);color:white;border:2px solid var(--accent);padding:12px 18px;cursor:pointer}button[aria-pressed=false]{background:white;color:var(--accent)}label{display:block;font-weight:600}input:not([type=checkbox]),select{width:100%;padding:12px;background:white;color:#17212B;border:1px solid #65747A;margin:8px 0 20px}input[type=checkbox]{width:20px;height:20px;vertical-align:middle}.actions{display:flex;flex-wrap:wrap;gap:12px;margin:20px 0}textarea{width:100%;resize:vertical}figure{margin:0;min-width:0}svg{width:100%;height:auto}[role=status]{min-height:1.6em}:focus-visible{outline:3px solid #B02E00;outline-offset:4px}blockquote{font:italic 1.25rem/1.6 "Libre Baskerville";margin:25px 0}small{font-size:14px}[hidden]{display:none!important}button:disabled{opacity:.65}@media(max-width:750px){main{padding:28px 22px}.spread,.notes{display:block}.note,.sheet,.ink{margin-bottom:22px}h1{font-size:2.65rem}.top{flex-wrap:wrap}th,td{padding:10px 8px}.stat{font-size:3rem}}@media print{.top{display:none}main{padding:0}.notes{break-inside:avoid}h1{font-size:32pt}footer{font-size:9pt}}"""


def table(d):
    return (
        '<div class="table-wrap" tabindex="0" role="region" aria-label="Data Table"><table><thead><tr>'
        + "".join('<th scope="col">' + H(x) + "</th>" for x in d["headers"])
        + "</tr></thead><tbody>"
        + "".join(
            "<tr>"
            + "".join(
                ('<th scope="row">' if i == 0 else "<td>")
                + H(x)
                + ("</th>" if i == 0 else "</td>")
                for i, x in enumerate(row)
            )
            + "</tr>"
            for row in d["rows"]
        )
        + "</tbody></table></div>"
    )


def web_document(d, c):
    concept, face, _, accent, paper, _, _ = c
    title = f'<header><p class="eyebrow">{H(d["tag"])}</p><h1>{H(d["title"])}</h1><p class="lead">{H(d["lead"])}</p></header>'
    metric = f'<aside class="measure"><strong class="stat">{H(d["stat"])}</strong><p>{H(d["label"])}</p></aside>'
    intro = f'<p class="intro">{H(d["intro"])}</p>'
    notes = [
        f'<section class="note"><h2>{H(t)}</h2><p>{H(b)}</p></section>'
        for t, b in d["sections"]
    ]
    ledger = table(d)
    ident = d["id"]
    css = ""
    if ident == "professional":
        body = (
            '<div class="dossier"><div>'
            + title
            + notes[0]
            + notes[1]
            + '</div><aside class="exceptions"><p class="eyebrow">RECONCILIATION / HOLD</p>'
            + metric
            + intro
            + notes[2]
            + '</aside><section class="wide"><h2>The Release Register</h2>'
            + ledger
            + "</section></div>"
        )
        css = ".dossier{display:grid;grid-template-columns:1.7fr 1fr;gap:45px}.exceptions{border-left:9px solid var(--accent);padding:28px;background:#FFE7DA}.dossier h1{font-size:clamp(2.4rem,4.5vw,4.4rem)}"
    elif ident == "legal":
        timeline = (
            '<section class="chronology">'
            + "".join(
                f'<article><p class="date">{H(row[0])}</p><h2>{H(row[1])}</h2><p>{H(row[2])}</p></article>'
                for row in d["rows"]
            )
            + "</section>"
        )
        body = (
            title
            + timeline
            + '<div class="spread"><div>'
            + intro
            + notes[0]
            + "</div><div>"
            + notes[1]
            + notes[2]
            + "</div></div><details><summary>Read the Exact Evidence Register</summary>"
            + ledger
            + "</details>"
        )
        css = '.chronology{display:grid;grid-template-columns:repeat(3,1fr);gap:30px;border-block:3px solid var(--accent);margin:35px 0;padding:15px 0}.date{font:700 3rem "Archivo";color:var(--accent);margin:0}.chronology h2{font-size:1.2rem}details{margin-top:25px}'
    elif ident == "business":
        body = (
            '<div class="proposal"><aside class="ink">'
            + metric
            + notes[2]
            + "</aside><div>"
            + title
            + intro
            + ledger
            + '</div><div class="wide spread">'
            + notes[0]
            + notes[1]
            + "</div></div>"
        )
        css = ".proposal{display:grid;grid-template-columns:1fr 2fr;gap:35px}.proposal h1{font-size:clamp(2.4rem,4.6vw,4.7rem)}.ink .note{border-color:white}.proposal .stat{font-size:3.4rem}"
    elif ident == "fun":
        body = (
            '<div class="poster"><div class="arrival">'
            + metric
            + notes[0]
            + "</div>"
            + title
            + '<div class="running"><div>'
            + ledger
            + "</div><div>"
            + notes[1]
            + notes[2]
            + intro
            + "</div></div></div>"
        )
        css = '.arrival{display:flex;align-items:start;gap:60px;border-bottom:2px solid;padding-bottom:22px}.poster header{padding:35px 0}.poster h1{font-size:clamp(3rem,8vw,8rem);max-width:15ch}.running{display:grid;grid-template-columns:1.2fr 1fr;gap:50px}.poster th{font-family:"Archivo"}.arrival .note{border:0;padding:0}'
    elif ident == "family":
        body = (
            title
            + '<section class="journey">'
            + "".join(
                f"<article><strong>{H(r[0])}</strong><p>{H(r[1])}</p><p>{H(r[2])}</p></article>"
                for r in d["rows"]
            )
            + '</section><div class="spread"><div class="sheet">'
            + notes[0]
            + "</div><div>"
            + notes[1]
            + notes[2]
            + "</div></div>"
            + intro
            + "<details><summary>Open the Full Weekend Register</summary>"
            + ledger
            + "</details>"
        )
        css = '.journey{display:flex;gap:20px;align-items:stretch;margin:30px 0}.journey article{flex:1;padding:22px;border-top:10px solid var(--accent);background:#FFF}.journey strong{font:2rem "Young Serif"}.journey article:nth-child(2){background:var(--accent);color:white}'
    elif ident == "presentation":
        body = (
            '<div class="argument"><div>'
            + title
            + intro
            + '</div><aside class="ink">'
            + metric
            + notes[1]
            + "</aside><section>"
            + notes[0]
            + "</section><section>"
            + notes[2]
            + '</section><div class="wide">'
            + ledger
            + "</div></div>"
        )
        css = ".argument{display:grid;grid-template-columns:1.2fr 1fr;gap:35px}.argument h1{font-size:clamp(2.5rem,4vw,4rem)}.argument .ink{border-radius:0 90px 0 0}.argument .ink .note{border-color:white}"
    elif ident == "school":
        bars = "".join(
            f'<text x="0" y="{28+i*70}">{H(row[0])} / {H(row[2])}</text><rect x="0" y="{40+i*70}" width="{int(row[2].split()[0])*2}" height="22" fill="currentColor"/>'
            for i, row in enumerate(d["rows"])
        )
        chart = (
            '<figure aria-label="Median Drainage Times" style="max-width:600px;margin-top:30px"><svg role="img" aria-labelledby="drainage-title drainage-desc" viewBox="0 0 405 255"><title id="drainage-title">Median Drainage Times</title><desc id="drainage-desc">Sand 40 seconds; garden soil 87 seconds; clay-rich soil 168 seconds. All bars begin at zero.</desc><g font-family="Work Sans,Arial" font-size="18" fill="currentColor">'
            + bars
            + '<path d="M0 38V222H360" fill="none" stroke="currentColor"/><text x="0" y="247">0</text><text x="112" y="247">60</text><text x="222" y="247">120</text><text x="327" y="247">180 sec</text></g></svg><figcaption>Median time to collect 100 mL; zero baseline. Raw trials remain in the table.</figcaption></figure>'
        )
        body = (
            '<div class="notebook">'
            + title
            + '<div class="lab-entry"><p class="eyebrow">OBSERVATION 01</p>'
            + intro
            + notes[0]
            + '</div><div class="results"><h2>The Trial Record</h2>'
            + ledger
            + chart
            + '</div><div class="spread">'
            + notes[1]
            + notes[2]
            + "</div>"
            + metric
            + "</div>"
        )
        css = ".notebook{max-width:920px;margin:auto;border-left:2px solid var(--accent);padding-left:45px}.lab-entry{border-block:1px solid var(--accent);padding:20px 0}.results{margin:30px 0;background:white;padding:24px}.notebook .measure{display:flex;align-items:center;gap:25px;margin-top:30px}"
    elif ident == "marketing":
        body = (
            title
            + '<div class="broadside"><div class="ink">'
            + metric
            + "</div><div>"
            + intro
            + notes[0]
            + '</div></div><section class="channels"><h2>Put the Message to Work</h2>'
            + ledger
            + '</section><div class="spread">'
            + notes[1]
            + notes[2]
            + "</div>"
        )
        css = "h1{font-size:clamp(3rem,8vw,8rem);max-width:17ch}.broadside{display:grid;grid-template-columns:1fr 2fr;gap:45px;margin:35px 0}.channels{margin-left:12%;border-left:10px solid var(--accent);padding-left:25px}"
    elif ident == "restaurant":
        body = (
            '<div class="service-time">'
            + metric
            + "<p>Late Summer / Supper Ledger</p></div>"
            + title
            + '<div class="courses-ledger">'
            + "".join(notes)
            + '</div><div class="spread"><div>'
            + intro
            + "</div><div>"
            + ledger
            + "</div></div>"
        )
        css = ".service-time{display:flex;justify-content:space-between;align-items:center;border-bottom:2px solid}.service-time .measure{display:flex;gap:20px;align-items:center}.service-time .stat{font-size:2rem}h1{font-size:clamp(2.6rem,4vw,4rem);max-width:25ch;margin-top:30px}.courses-ledger{display:grid;grid-template-columns:repeat(3,1fr);gap:35px;margin:40px 0}.courses-ledger .note{border-top:8px solid var(--accent)}"
    else:
        body = (
            '<aside class="stop"><strong>STOP / Preserve the Evidence</strong><p>'
            + H(d["closing"])
            + "</p></aside>"
            + title
            + '<div class="runbook"><section>'
            + "".join(notes)
            + '</section><aside class="sheet">'
            + metric
            + intro
            + "<h2>Verification Register</h2>"
            + ledger
            + "</aside></div>"
        )
        css = ".stop{background:#9E202C;color:white;padding:20px 28px;margin-bottom:35px}.runbook{display:grid;grid-template-columns:1fr 1.3fr;gap:35px}.runbook .note{border-left:7px solid var(--accent);padding:20px}.runbook .stat{font-size:3rem}h1{font-size:clamp(2.7rem,4.5vw,4.6rem)}"
    css += "@media(max-width:750px){.dossier,.proposal,.running,.argument,.broadside,.runbook,.chronology,.courses-ledger{display:block}.chronology article{margin:24px 0}.arrival,.journey,.service-time,.service-time .measure{display:block}.journey article{margin:18px 0}.exceptions{margin:25px 0}.notebook{padding-left:20px}.channels{margin:20px 0;padding-left:15px}.poster h1{font-size:3.3rem}}"
    return body + f'<p class="closing">{H(d["closing"])}</p>', css


def shade(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(element)


def native(d, c):
    _, _, face, accent, paper, orientation, _ = c
    doc = Document()
    sec = doc.sections[0]
    landscape = orientation == "landscape"
    sec.page_width = Inches(11.7 if landscape else 8.3)
    sec.page_height = Inches(8.3 if landscape else 11.7)
    sec.top_margin = sec.bottom_margin = Inches(0.48)
    sec.left_margin = sec.right_margin = Inches(0.55)
    for element in list(doc.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        st = doc.styles[name]
        st.font.name = face
        st.font.size = Pt(10.5)
        st.font.color.rgb = RGBColor.from_string("17212B")
        st.paragraph_format.space_after = Pt(6)
        for e in list(st.element.iter(qn("w:pBdr"))):
            e.getparent().remove(e)

    def p(where, text, size=10.5, bold=False, color="#17212B", style=None):
        el = where.add_paragraph(style=style)
        el.paragraph_format.space_after = Pt(7)
        el.paragraph_format.line_spacing = 1.08
        if len(text) > 80 and size <= 12 and not bold:
            width = getattr(where, "width", None) or (
                sec.page_width - sec.left_margin - sec.right_margin
            )
            el.paragraph_format.right_indent = Pt(max(0, width / 12700 - size * 29))
        r = el.add_run(text)
        r.font.name = face
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color[1:])
        return el

    def note(where, item):
        p(where, item[0], 14, True, accent, "Heading 1")
        p(where, item[1])

    def grid(where, widths):
        t = where.add_table(rows=1, cols=len(widths))
        t.autofit = False
        for col, w in zip(t.columns, widths):
            col.width = Inches(w)
        for cell, w in zip(t.rows[0].cells, widths):
            cell.width = Inches(w)
        return t.rows[0].cells

    def ledger(where):
        t = where.add_table(rows=1, cols=len(d["headers"]))
        if d["id"] == "technical":
            t.autofit = False
            for col, w in zip(t.columns, [1.45, 1.05, 1.05]):
                col.width = Inches(w)
            for cell, w in zip(t.rows[0].cells, [1.45, 1.05, 1.05]):
                cell.width = Inches(w)
        for i, x in enumerate(d["headers"]):
            cell = t.cell(0, i)
            shade(cell, accent)
            p(cell, x, 9, True, "#FFFFFF")
        flag = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(flag)
        for row in d["rows"]:
            for cell, x in zip(t.add_row().cells, row):
                shade(cell, "#FFFFFF")
                p(cell, x, 9)
        for row in t.rows:
            flag = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(flag)

    def title(where):
        p(where, d["title"], 27 if landscape else 30, True, "#17212B", "Title")
        p(where, d["lead"], 12, True, accent)

    def metric(where):
        p(where, d["stat"], 29, True, accent)
        p(where, d["label"], 10)

    p(doc, d["tag"], 9, True, accent)
    ident = d["id"]
    if ident == "professional":
        title(doc)
        left, right = grid(doc, [4.5, 2.5])
        shade(right, "#FFE7DA")
        note(left, d["sections"][0])
        note(left, d["sections"][1])
        metric(right)
        p(right, d["intro"], 10)
        note(doc, d["sections"][2])
        ledger(doc)
    elif ident == "legal":
        title(doc)
        cells = grid(doc, [3.5] * 3)
        for cell, row in zip(cells, d["rows"]):
            shade(cell, "#E4EDFF")
            p(cell, row[0], 26, True, accent)
            p(cell, row[1], 12, True)
            p(cell, row[2])
        p(doc, d["intro"])
        cells = grid(doc, [3.5] * 3)
        for cell, item in zip(cells, d["sections"]):
            note(cell, item)
    elif ident == "business":
        left, right = grid(doc, [3.3, 7.1])
        shade(left, paper)
        metric(left)
        note(left, d["sections"][2])
        title(right)
        p(right, d["intro"])
        ledger(right)
        cells = grid(doc, [5.2, 5.2])
        note(cells[0], d["sections"][0])
        note(cells[1], d["sections"][1])
    elif ident == "fun":
        metric(doc)
        title(doc)
        ledger(doc)
        p(doc, d["intro"])
        cells = grid(doc, [2.3] * 3)
        for cell, item in zip(cells, d["sections"]):
            note(cell, item)
    elif ident == "family":
        title(doc)
        cells = grid(doc, [3.5] * 3)
        for cell, row in zip(cells, d["rows"]):
            shade(cell, "#FFFFFF")
            p(cell, row[0], 22, True, accent)
            p(cell, row[1], 11, True)
            p(cell, row[2])
        p(doc, d["intro"])
        cells = grid(doc, [3.5] * 3)
        for cell, item in zip(cells, d["sections"]):
            note(cell, item)
    elif ident == "presentation":
        left, right = grid(doc, [5.8, 4.6])
        title(left)
        p(left, d["intro"])
        shade(right, "#FFFFFF")
        metric(right)
        note(right, d["sections"][1])
        cells = grid(doc, [5.2] * 2)
        note(cells[0], d["sections"][0])
        note(cells[1], d["sections"][2])
        ledger(doc)
    elif ident == "school":
        title(doc)
        p(doc, "OBSERVATION 01", 10, True, accent)
        p(doc, d["intro"])
        note(doc, d["sections"][0])
        ledger(doc)
        cells = grid(doc, [3.5] * 2)
        note(cells[0], d["sections"][1])
        note(cells[1], d["sections"][2])
        p(doc, d["stat"] + " / " + d["label"], 14, True, accent)
    elif ident == "marketing":
        title(doc)
        left, right = grid(doc, [2.2, 4.8])
        shade(left, paper)
        metric(left)
        p(right, d["intro"])
        note(right, d["sections"][0])
        ledger(doc)
        cells = grid(doc, [3.5] * 2)
        note(cells[0], d["sections"][1])
        note(cells[1], d["sections"][2])
    elif ident == "restaurant":
        p(doc, d["stat"] + " / " + d["label"], 11, True, accent)
        title(doc)
        cells = grid(doc, [3.5] * 3)
        for cell, item in zip(cells, d["sections"]):
            note(cell, item)
        p(doc, d["intro"])
        ledger(doc)
    else:
        p(doc, "STOP / Preserve the Evidence", 13, True, "#9E202C")
        title(doc)
        left, right = grid(doc, [3.4, 3.6])
        for item in d["sections"]:
            note(left, item)
        shade(right, "#FFFFFF")
        metric(right)
        p(right, d["intro"], 10)
        ledger(right)
    p(doc, d["closing"], 9, True, accent)
    p(
        sec.footer,
        "Fictional Dazzler example / Apache-2.0 / Editable text and tables",
        8,
    )
    # Native page surface, independent of the host's print-background preference.
    from lxml import etree

    sec.header.paragraphs[0].add_run()._r.append(
        etree.fromstring(
            f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:rect id="DazzlerPageSurface" style="position:absolute;margin-left:0;margin-top:0;width:{sec.page_width/12700}pt;height:{sec.page_height/12700}pt;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" fillcolor="{paper}" stroked="f"><v:fill color="{paper}"/></v:rect></w:pict>'
        )
    )
    doc.save(O / "docx" / f"{ident}.docx")


def documents():
    briefs = []
    for ident, c in D.items():
        d = json.loads((OLD / "data" / f"{ident}.json").read_text(encoding="utf-8"))
        concept, face, nativeface, accent, paper, orientation, why = c
        prompt = f'Use Dazzler to create a fictional {ident} document and web edition for this task: {d["lead"]} Retain every supplied fact and qualification. {why} Use {face} for the web and verified {nativeface} for native output, with {accent} against {paper}. Author a distinct composition rather than recoloring the previous release. Keep native text and tables editable and verify narrow layouts.'
        d.update(
            font=face,
            word=nativeface,
            accent=accent,
            paper=paper,
            layout=concept,
            prompt=prompt,
        )
        d["fontSelection"] = {
            "id": face.lower().replace(" ", "-"),
            "css_family": face,
            "license": "OFL-1.1",
            "coverage_checked": False,
        }
        d["design"].update(
            accent=accent, secondary=accent, paper=paper, accentText=accent
        )
        save(O / "data" / f"{ident}.json", d)
        body, css = web_document(d, c)
        page = (
            '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'
            + H(d["title"])
            + '</title><link rel="stylesheet" href="../fonts/fonts.css"><style>:root{--accent:'
            + accent
            + ";--paper:"
            + paper
            + ';--display:"'
            + face
            + '"}'
            + BASE
            + css
            + '</style><body><nav class="top"><a href="../index.html">Dazzler Design Gallery</a><span>Fictional '
            + H(concept)
            + "</span></nav><main>"
            + body
            + "</main><footer>Original Dazzler design and synthetic facts. Apache-2.0. Font notices accompany the files.</footer></body></html>"
        )
        (O / "html" / f"{ident}.html").write_text(titled(page), encoding="utf-8")
        native(d, c)
        briefs.append(
            dict(
                id=ident,
                prompt=prompt,
                direction=concept,
                rationale=why,
                font=face,
                word=nativeface,
                surface=paper + " / " + accent,
                orientation=orientation,
                composition=concept,
                directionsConsidered=[
                    "A continuous annotated essay",
                    concept,
                    "An image-led cover with appended evidence",
                ],
                selectionReason=why,
                structuralChanges=[
                    why,
                    "Reorders the factual register and prose into the new "
                    + concept.lower()
                    + " reading sequence.",
                    "Replaces the previous display hierarchy with "
                    + face
                    + " and "
                    + orientation
                    + " document geometry.",
                ],
            )
        )
    save(S / "document-briefs.json", briefs)


def start():
    S.mkdir(parents=True, exist_ok=True)
    if not OLD.exists():
        shutil.copytree(R / "skills/dazzler-frontend/assets/templates", OLD)
    if not O.exists():
        shutil.copytree(OLD, O)
    previous = S / "previous-v0.29.0"
    previous.mkdir(exist_ok=True)
    for p in (OLD / "previews").glob("*.jpg"):
        shutil.copy2(p, previous / p.name)


if __name__ == "__main__":
    start()
    documents()
